import io
import os
import sys
import textwrap
import streamlit as st
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import matplotlib.patches as patches
import plotly.express as px
import plotly.graph_objects as go
from plotly.subplots import make_subplots

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Ensure project root directory is on sys.path for robust sub-module imports
ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

from config import (
    MAINTENANCE_CAUSES,
    parse_machine_size,
    resolve_section_config,
    resolve_npt_category,
    NPT_CATEGORIES,
    EXCEL_SIZES,
)

DEFAULT_TOP_10_CAUSES = [
    "No Demand",
    "Manpower Short*",
    "Mold Problem*",
    "Machine Problem*",
    "Power Breakdown (Unscheduled)*",
    "Robot Problem*",
    "Sample + Mold Test (RND)*",
    "Product Jam*",
    "Mold Change*",
    "RMCS Problem*",
]


def get_col(df, candidates, default=None):
    for c in candidates:
        if c in df.columns:
            return c
    return default


# =========================================================
# 0. COMPANY STANDARD EXCEL STYLER (.XLSX EXPORT)
# =========================================================
def convert_df_to_styled_excel(df, sheet_name="NPT_Report"):
    """Generates Excel with company yellow headers, thin borders, and bold red subtotals."""
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    header_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="000000")
    data_font = Font(name="Calibri", size=11, bold=False, color="000000")
    total_font = Font(name="Calibri", size=11, bold=True, color="FF0000")

    thin_border_side = Side(border_style="thin", color="BFBFBF")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    double_bottom_side = Side(border_style="double", color="000000")
    total_top_side = Side(border_style="thin", color="000000")
    total_border = Border(left=thin_border_side, right=thin_border_side, top=total_top_side, bottom=double_bottom_side)

    for col_num, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_num, value=str(col_name))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border

    for row_idx, row in df.iterrows():
        excel_row_num = row_idx + 2
        is_total_row = any("sub total" in str(v).lower() or "total" in str(v).lower() or "summary >>" in str(v).lower() for v in row.values)

        for col_idx, (col_name, val) in enumerate(row.items(), 1):
            cell = ws.cell(row=excel_row_num, column=col_idx)
            col_lower = str(col_name).lower()
            val_str = str(val).strip()

            if is_total_row:
                cell.font = total_font
                cell.border = total_border
            else:
                cell.font = data_font
                cell.border = cell_border

            if isinstance(val, (int, float)) and pd.notna(val):
                cell.value = val
                if any(k in col_lower for k in ["qty", "nos", "count", "rank"]):
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif any(k in col_lower for k in ["%", "share", "pct", "deviation", "impact"]):
                    cell.number_format = '0.00%'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                clean_num_str = val_str.replace(",", "").replace("%", "")
                try:
                    num_val = float(clean_num_str)
                    if "%" in val_str:
                        cell.value = num_val / 100.0
                        cell.number_format = '0.00%'
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    elif clean_num_str.isdigit():
                        cell.value = int(clean_num_str)
                        cell.number_format = '#,##0'
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.value = num_val
                        cell.number_format = '#,##0.00'
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                except ValueError:
                    cell.value = val_str
                    if is_total_row:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif col_lower in ["date", "size", "mc size", "rank", "mc sl", "status"]:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

    wb.save(output)
    return output.getvalue()


# =========================================================
# 1. PARSING ENGINES (8 AM TO 8 AM PRODUCTION SLICING)
# =========================================================
@st.cache_data
def m3_parse_downtime_workbook(file_bytes):
    file_stream = io.BytesIO(file_bytes)
    xls = pd.ExcelFile(file_stream)
    sheet_name = (
        "NPT Details"
        if "NPT Details" in xls.sheet_names
        else ("DowntimeReport" if "DowntimeReport" in xls.sheet_names else xls.sheet_names[0])
    )
    df_raw = pd.read_excel(xls, sheet_name=sheet_name, header=None)

    header_idx = None
    for idx, row in df_raw.iterrows():
        row_str = " ".join([str(v) for v in row.values])
        if "Machine" in row_str and ("Duration" in row_str or "Cause" in row_str):
            header_idx = idx
            break

    df_clean = (
        pd.read_excel(xls, sheet_name=sheet_name, skiprows=header_idx)
        if header_idx is not None
        else pd.read_excel(xls, sheet_name=sheet_name)
    )
    df_clean.columns = [str(c).strip() for c in df_clean.columns]

    to_col = get_col(df_clean, ["To Time", "ToTime", "End Time", "EndTime"])
    from_col = get_col(df_clean, ["From Time", "FromTime", "Start Time", "StartTime"])
    cause_col = get_col(df_clean, ["Cause", "Causes", "Reason", "Defect"], "Cause")
    mc_col = get_col(df_clean, ["Machine", "MC SL", "Smart Manu"], "Machine")

    df_clean["From_DT"] = pd.to_datetime(df_clean[from_col], errors="coerce")
    df_clean["To_DT"] = pd.to_datetime(df_clean[to_col], errors="coerce") if to_col else pd.NaT
    df_clean["Is_Ongoing"] = df_clean["To_DT"].isna()

    # Prevent data loss when cause has not been assigned yet
    raw_cause = df_clean[cause_col].fillna("Unassigned / Pending Log*").astype(str)
    df_clean["Cause"] = raw_cause
    df_clean["CauseClean"] = raw_cause.str.replace("*", "", regex=False).str.strip()
    df_clean["Category"] = df_clean["Cause"].apply(resolve_npt_category)
    df_clean["Is_Maintenance"] = df_clean["Cause"].isin(MAINTENANCE_CAUSES)
    df_clean["Is_SMED"] = df_clean["CauseClean"].str.lower() == "mold change"

    # Auto-resolve section configuration dynamically
    sec_name, total_mcs, daily_avail_hrs, pos_map, size_counts, unique_sizes = resolve_section_config(df_clean)

    df_clean["Detected_Section"] = sec_name
    df_clean["Position"] = df_clean[mc_col].astype(str).map(pos_map).fillna(df_clean[mc_col].astype(str))
    df_clean["Size"] = df_clean[mc_col].apply(parse_machine_size)

    # Floor Detection for Multi-Floor Plants (Plastic-3 FF: 90/120/160; GF: 250-800)
    def determine_floor(row):
        pos = str(row["Position"]).strip().upper()
        if pos.startswith(("A", "B", "C", "D")) and any(z in pos for z in ["-90", "-120", "-160"]):
            return "FF"
        return "GF"

    df_clean["Floor"] = df_clean.apply(determine_floor, axis=1)

    return df_clean


def slice_downtime_by_operational_days(df_parsed, cutoff_cutoff_dt=None):
    """
    Slices multi-day and overnight stoppages into 8 AM to 8 AM operational days.
    Operational day D = window from (D 08:00:00) to (D+1 08:00:00).
    Ongoing incidents are evaluated up to cutoff_cutoff_dt.
    """
    records = []

    for _, row in df_parsed.iterrows():
        f_val = row["From_DT"]
        t_val = row["To_DT"]
        is_ongoing = row["Is_Ongoing"]

        if pd.isna(f_val):
            continue

        if is_ongoing:
            if cutoff_cutoff_dt is not None:
                t_eval = min(pd.to_datetime(cutoff_cutoff_dt), pd.Timestamp.now())
            else:
                t_eval = pd.Timestamp.now()
            if t_eval < f_val:
                t_eval = f_val
        else:
            t_eval = t_val

        op_start = (f_val - pd.Timedelta(hours=8)).normalize()
        op_end = (t_eval - pd.Timedelta(hours=8)).normalize()

        if pd.isna(op_start) or pd.isna(op_end):
            continue

        dates = pd.date_range(op_start, op_end, freq="D")
        for d in dates:
            d_win_start = d + pd.Timedelta(hours=8)
            d_win_end = d + pd.Timedelta(days=1, hours=8)

            s_start = max(f_val, d_win_start)
            s_end = min(t_eval, d_win_end)

            slice_hrs = (s_end - s_start).total_seconds() / 3600.0
            if slice_hrs >= 0:
                r_dict = row.to_dict()
                r_dict["DateClean"] = d
                r_dict["DateStr"] = d.strftime("%Y-%m-%d")
                r_dict["MonthName"] = d.strftime("%B")
                r_dict["DayNum"] = d.day
                r_dict["YearMonth"] = d.to_period("M")
                r_dict["Hours"] = slice_hrs
                r_dict["Slice_Ongoing"] = is_ongoing and (d == op_end)
                records.append(r_dict)

    df_res = pd.DataFrame(records)
    return df_res


# =========================================================
# 2. COMPUTATION HELPERS
# =========================================================
def m3_compute_size_wise_npt(df_scope, cutoff_days, total_plant_mcs, size_nos_map, unique_sizes):
    records = []
    tot_hrs_all = df_scope["Hours"].sum() if not df_scope.empty else 0.0

    for sz in unique_sizes:
        nos = size_nos_map.get(sz, 0)
        sz_hrs = df_scope[df_scope["Size"] == sz]["Hours"].sum() if not df_scope.empty else 0.0
        avail_hrs = nos * 24.0 * max(1, cutoff_days)
        cap_pct = (sz_hrs / avail_hrs * 100.0) if avail_hrs > 0 else 0.0

        records.append({
            "MC Size": sz,
            "Nos": nos,
            "NPT hrs": round(sz_hrs, 1),
            "NPT %": round(cap_pct, 1),
        })

    df_res = pd.DataFrame(records)
    tot_avail = total_plant_mcs * 24.0 * max(1, cutoff_days)
    summary_pct = (tot_hrs_all / tot_avail * 100.0) if tot_avail > 0 else 0.0

    return df_res, tot_hrs_all, summary_pct


def m3_compute_size_wise_mom(df_curr_mtd, df_prev_mtd, cutoff_days, total_plant_mcs, size_nos_map, unique_sizes, prev_abbr="Aug", curr_abbr="Sep", include_variances=False):
    """Computes side-by-side MoM capacity loss by machine size."""
    records = []
    c_curr_hrs = f"{curr_abbr} NPT (Hrs)"
    c_prev_hrs = f"{prev_abbr} NPT (Hrs)"
    c_curr_pct = f"{curr_abbr} NPT %"
    c_prev_pct = f"{prev_abbr} NPT %"

    for sz in unique_sizes:
        nos = size_nos_map.get(sz, 0)
        c_hrs = df_curr_mtd[df_curr_mtd["Size"] == sz]["Hours"].sum() if not df_curr_mtd.empty else 0.0
        p_hrs = df_prev_mtd[df_prev_mtd["Size"] == sz]["Hours"].sum() if not df_prev_mtd.empty else 0.0
        avail_hrs = nos * 24.0 * max(1, cutoff_days)

        c_cap_pct = (c_hrs / avail_hrs * 100.0) if avail_hrs > 0 else 0.0
        p_cap_pct = (p_hrs / avail_hrs * 100.0) if avail_hrs > 0 else 0.0
        diff_hrs = c_hrs - p_hrs
        diff_pct = c_cap_pct - p_cap_pct

        row = {
            "MC Size": sz,
            "Nos": nos,
            c_curr_hrs: round(c_hrs, 1),
            c_prev_hrs: round(p_hrs, 1),
            c_curr_pct: round(c_cap_pct, 1),
            c_prev_pct: round(p_cap_pct, 1),
        }
        if include_variances:
            row["Variance (Hrs)"] = round(diff_hrs, 1)
            row["Variance (NPT %)"] = round(diff_pct, 1)
        records.append(row)

    return pd.DataFrame(records)


def m3_compute_smed_table(df_scope, max_days=7):
    smed_df = df_scope[df_scope["Is_SMED"]].copy()
    if smed_df.empty:
        return pd.DataFrame(), 0, 0.0, 0.0

    daily = (
        smed_df.groupby("DateClean")
        .agg(
            Mold_Change_Qty=("Hours", "count"),
            Total_Time=("Hours", "sum"),
            Involved_Mcs=(
                "Position",
                lambda x: ", ".join(sorted(set(str(v) for v in x if str(v) != "-"))),
            ),
        )
        .reset_index()
        .sort_values("DateClean")
    )
    daily["Avg_SMED"] = (daily["Total_Time"] / daily["Mold_Change_Qty"] * 60.0).round(2)
    daily["Date"] = daily["DateClean"].dt.strftime("%-d-%b")
    daily["Total_Time"] = daily["Total_Time"].round(2)

    tot_mtd_setups = int(daily["Mold_Change_Qty"].sum())
    tot_mtd_time = float(daily["Total_Time"].sum())
    mtd_avg_smed = (tot_mtd_time / tot_mtd_setups * 60.0) if tot_mtd_setups > 0 else 0.0

    daily_display = daily.tail(max_days).reset_index(drop=True) if len(daily) > max_days else daily.reset_index(drop=True)
    return daily_display[["Date", "Mold_Change_Qty", "Total_Time", "Avg_SMED", "Involved_Mcs"]], tot_mtd_setups, tot_mtd_time, mtd_avg_smed


def m3_compute_maint_daily_table(df_scope, cutoff_days, daily_available_hrs, total_plant_mcs, max_days=7):
    maint_causes = [
        "Machine Problem*",
        "Robot Problem*",
        "Controller Problem*",
        "RMCS Problem*",
        "Oil or water Leakage*",
    ]
    all_dates = sorted(df_scope["DateClean"].dropna().unique())
    dates_display = all_dates[-max_days:] if len(all_dates) > max_days else all_dates

    records = []
    for dt in dates_display:
        dt_df = df_scope[df_scope["DateClean"] == dt]
        row = {"Date": dt.strftime("%-d-%b")}

        maint_sum = 0.0
        for mc_cause in maint_causes:
            c_hrs = dt_df[dt_df["Cause"] == mc_cause]["Hours"].sum()
            row[mc_cause] = round(c_hrs, 2)
            maint_sum += c_hrs

        row["Total Hours"] = round(maint_sum, 2)
        share_pct = (maint_sum / daily_available_hrs * 100.0) if daily_available_hrs > 0 else 0.0
        row["Total Share"] = f"{round(share_pct):.0f}%"
        records.append(row)

    df_display = pd.DataFrame(records)

    mtd_summary_row = {"Date": f"Total (1-{cutoff_days}) >>"}
    tot_maint_all = 0.0
    for mc_cause in maint_causes:
        c_tot = df_scope[df_scope["Cause"] == mc_cause]["Hours"].sum()
        mtd_summary_row[mc_cause] = round(c_tot, 1)
        tot_maint_all += c_tot

    mtd_summary_row["Total Hours"] = round(tot_maint_all, 1)
    tot_avail_period = total_plant_mcs * 24.0 * max(1, cutoff_days)
    mtd_share = (tot_maint_all / tot_avail_period * 100.0) if tot_avail_period > 0 else 0.0
    mtd_summary_row["Total Share"] = f"{round(mtd_share):.0f}%"

    return df_display, mtd_summary_row


def m3_compute_consolidated_daily_log(df_day):
    if df_day.empty:
        return pd.DataFrame()

    records = []
    for (pos, mc), grp in df_day.groupby(["Position", "Machine"]):
        has_ongoing = grp["Slice_Ongoing"].any() or grp["Is_Ongoing"].any()
        causes = ", ".join(
            sorted(
                set(
                    str(c).replace("*", "").strip()
                    for c in grp["Cause"].dropna().unique()
                )
            )
        )
        cats = ", ".join(sorted(set(str(c).strip() for c in grp["Category"].dropna().unique())))
        day_hrs = min(24.0, grp["Hours"].sum())  # Strict 24h single-machine physical cap
        start_t = str(grp["From_DT"].min())[:16]

        if has_ongoing:
            status = "Active / In Progress (Active)"
            dur_display = f"{day_hrs:.2f} h (Active)"
        else:
            status = "Closed"
            dur_display = f"{day_hrs:.2f} h"

        records.append({
            "Position": pos,
            "Machine": mc,
            "Department": cats,
            "Status": status,
            "Combined Causes": causes,
            "Closed Hours": round(day_hrs, 2),
            "Duration": dur_display,
            "Start Time": start_t,
            "Is_Ongoing": has_ongoing,
        })

    res = (
        pd.DataFrame(records)
        .sort_values(by=["Is_Ongoing", "Closed Hours"], ascending=[False, False])
        .reset_index(drop=True)
    )
    return res


# =========================================================
# 3. EXECUTIVE 4-GRID 1-PAGE REPORT (PROFESSIONAL 2×2)
# =========================================================
def m3_generate_2x2_executive_jpg(
    sel_date_obj,
    cutoff_day,
    top_10_causes,
    curr_share_dict,
    prev_share_dict,
    curr_hours_dict,
    tot_curr_mtd_hrs,
    tot_prev_mtd_hrs,
    curr_month_name,
    prev_month_name,
    df_size_grid,
    size_tot_hrs,
    size_summary_pct,
    df_smed_grid,
    smed_tot_qty,
    smed_tot_time,
    smed_avg_min,
    df_maint_grid,
    maint_summary_dict,
    section_label,
    total_plant_mcs,
):
    fig, ax = plt.subplots(figsize=(20.0, 11.5), dpi=240)
    fig.patch.set_facecolor("#f8fafc")
    ax.set_facecolor("#f8fafc")
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis("off")

    curr_abbr = curr_month_name[:3].capitalize()
    prev_abbr = prev_month_name[:3].capitalize()
    start_date_str = f"{curr_abbr} 01"
    end_date_str = f"{curr_abbr} {cutoff_day:02d}, {sel_date_obj.year}"
    span_title = f"{start_date_str.upper()} TO {end_date_str.upper()} NPT ANALYSIS"

    banner = patches.FancyBboxPatch(
        (1.5, 92.5), 97.0, 6.5, boxstyle="round,pad=0.2,rounding_size=0.4", facecolor="#091e3a", edgecolor="none"
    )
    ax.add_patch(banner)
    ax.text(3.5, 96.8, "OPERATIONAL ANALYTICS — NON-PRODUCTIVE TIME (NPT) REPORT", color="#ffffff", fontsize=14.5, fontweight="bold", va="center")
    ax.text(3.5, 94.2, f"{span_title}  |  {section_label} ({total_plant_mcs} Machines Baseline | 8 AM–8 AM Production Day)", color="#38bdf8", fontsize=9.5, fontweight="bold", va="center")
    ax.text(96.5, 95.5, f"Report Date: {sel_date_obj.strftime('%d-%m-%Y')}", color="#94a3b8", fontsize=9.0, ha="right", va="center")

    w_box, h_box = 47.8, 43.5
    p1 = patches.FancyBboxPatch((1.5, 47.5), w_box, h_box, boxstyle="round,pad=0.2,rounding_size=0.5", facecolor="#ffffff", edgecolor="#cbd5e1", linewidth=0.8)
    p2 = patches.FancyBboxPatch((50.7, 47.5), w_box, h_box, boxstyle="round,pad=0.2,rounding_size=0.5", facecolor="#ffffff", edgecolor="#cbd5e1", linewidth=0.8)
    p3 = patches.FancyBboxPatch((1.5, 2.5), w_box, h_box, boxstyle="round,pad=0.2,rounding_size=0.5", facecolor="#ffffff", edgecolor="#cbd5e1", linewidth=0.8)
    p4 = patches.FancyBboxPatch((50.7, 2.5), w_box, h_box, boxstyle="round,pad=0.2,rounding_size=0.5", facecolor="#ffffff", edgecolor="#cbd5e1", linewidth=0.8)

    for p in [p1, p2, p3, p4]:
        ax.add_patch(p)

    hdr_g1 = patches.Rectangle((1.5, 87.8), w_box, 3.2, facecolor="#091e3a", edgecolor="none")
    ax.add_patch(hdr_g1)
    ax.text(3.0, 89.4, f"TOP 10 NPT SHARE IMPACT ({curr_abbr} 01–{cutoff_day:02d} vs {prev_abbr} 01–{cutoff_day:02d})", color="#ffffff", fontsize=8.8, fontweight="bold", va="center")
    
    ax.add_patch(patches.Rectangle((33.0, 88.8), 1.8, 1.1, facecolor="#ef4444", edgecolor="none"))
    ax.text(35.2, 89.4, f"{curr_abbr}", color="#ffffff", fontsize=7.6, va="center")
    ax.add_patch(patches.Rectangle((40.5, 88.8), 1.8, 1.1, facecolor="#94a3b8", edgecolor="none"))
    ax.text(42.7, 89.4, f"{prev_abbr}", color="#ffffff", fontsize=7.6, va="center")

    ax.add_patch(patches.Rectangle((1.5, 85.0), w_box, 2.6, facecolor="#f1f5f9", edgecolor="#e2e8f0", linewidth=0.5))
    ax.text(3.0, 86.3, "Cause Description", color="#0f172a", fontsize=7.4, fontweight="bold", va="center")
    ax.text(17.5, 86.3, f"{curr_abbr} (Hrs)", color="#0f172a", fontsize=7.4, fontweight="bold", ha="right", va="center")
    ax.text(19.5, 86.3, f"Share % ({curr_abbr} vs {prev_abbr})", color="#0f172a", fontsize=7.4, fontweight="bold", va="center")
    ax.text(46.0, 86.3, "Variance", color="#0f172a", fontsize=7.4, fontweight="bold", ha="center", va="center")

    hdr_g2 = patches.Rectangle((50.7, 87.8), w_box, 3.2, facecolor="#091e3a", edgecolor="none")
    ax.add_patch(hdr_g2)
    ax.text(53.0, 89.4, f"MC SIZE-WISE NPT CAPACITY LOSS ({curr_abbr} 01–{cutoff_day:02d}, {sel_date_obj.year})", color="#ffffff", fontsize=8.8, fontweight="bold", va="center")

    ax.add_patch(patches.Rectangle((50.7, 85.0), w_box, 2.6, facecolor="#f1f5f9", edgecolor="#e2e8f0", linewidth=0.5))
    ax.text(53.5, 86.3, "Mc Size", color="#0f172a", fontsize=7.6, fontweight="bold", va="center")
    ax.text(65.0, 86.3, "Nos", color="#0f172a", fontsize=7.6, fontweight="bold", va="center")
    ax.text(77.0, 86.3, "NPT hrs", color="#0f172a", fontsize=7.6, fontweight="bold", va="center")
    ax.text(90.0, 86.3, "NPT %", color="#0f172a", fontsize=7.6, fontweight="bold", va="center")

    hdr_g3 = patches.Rectangle((1.5, 42.8), w_box, 3.2, facecolor="#091e3a", edgecolor="none")
    ax.add_patch(hdr_g3)
    ax.text(3.0, 44.4, "SMED", color="#ffffff", fontsize=9.0, fontweight="bold", va="center")

    ax.add_patch(patches.Rectangle((1.5, 40.2), w_box, 2.4, facecolor="#f1f5f9", edgecolor="#e2e8f0", linewidth=0.5))
    ax.text(2.8, 41.4, "Date", color="#0f172a", fontsize=7.2, fontweight="bold", va="center")
    ax.text(8.8, 41.4, "Mold Qty", color="#0f172a", fontsize=7.2, fontweight="bold", va="center")
    ax.text(16.5, 41.4, "Total (Hr)", color="#0f172a", fontsize=7.2, fontweight="bold", va="center")
    ax.text(24.5, 41.4, "Avg SMED", color="#0f172a", fontsize=7.2, fontweight="bold", va="center")
    ax.text(34.0, 41.4, "Involved Mcs", color="#0f172a", fontsize=7.2, fontweight="bold", va="center")

    hdr_g4 = patches.Rectangle((50.7, 42.8), w_box, 3.2, facecolor="#00a859", edgecolor="none")
    ax.add_patch(hdr_g4)
    ax.text(53.0, 44.4, "MC MAINTENANCE RELATED ISSUE", color="#ffffff", fontsize=9.0, fontweight="bold", va="center")

    x_g4 = [52.2, 58.0, 64.5, 71.8, 78.8, 85.5, 91.5, 95.8]
    ax.add_patch(patches.Rectangle((50.7, 40.2), w_box, 2.4, facecolor="#f1f5f9", edgecolor="#e2e8f0", linewidth=0.5))
    ax.text(x_g4[0], 41.4, "Date", color="#0f172a", fontsize=7.0, fontweight="bold", va="center")
    ax.text(x_g4[1], 41.4, "Machine*", color="#0f172a", fontsize=7.0, fontweight="bold", va="center")
    ax.text(x_g4[2], 41.4, "Robot*", color="#0f172a", fontsize=7.0, fontweight="bold", va="center")
    ax.text(x_g4[3], 41.4, "Controller*", color="#0f172a", fontsize=7.0, fontweight="bold", va="center")
    ax.text(x_g4[4], 41.4, "RMCS*", color="#0f172a", fontsize=7.0, fontweight="bold", va="center")
    ax.text(x_g4[5], 41.4, "Oil/Water*", color="#0f172a", fontsize=7.0, fontweight="bold", va="center")
    ax.text(x_g4[6], 41.4, "Total Hrs", color="#0f172a", fontsize=7.0, fontweight="bold", va="center")
    ax.text(x_g4[7], 41.4, "Share", color="#0f172a", fontsize=7.0, fontweight="bold", va="center")

    y_g1 = 82.8
    y_step_g1 = 3.25
    all_shares = [curr_share_dict.get(c, 0.0) for c in top_10_causes] + [prev_share_dict.get(c, 0.0) for c in top_10_causes]
    max_share = max(all_shares + [35.0])

    for c in top_10_causes[:10]:
        val_curr = curr_share_dict.get(c, 0.0)
        val_prev = prev_share_dict.get(c, 0.0)
        hrs_curr = curr_hours_dict.get(c, 0.0)
        diff = val_curr - val_prev

        c_label = c.replace("*", "")[:20]
        ax.text(3.0, y_g1 - 0.2, c_label, color="#b91c1c" if "Problem" in c else "#0f172a", fontsize=7.2, fontweight="bold", va="center")
        ax.text(17.5, y_g1 - 0.2, f"{hrs_curr:.1f}h", color="#0f172a", fontsize=7.0, fontweight="bold", ha="right", va="center")

        bar_x = 19.5
        bar_max_w = 21.0
        w_curr = (val_curr / max_share) * bar_max_w
        w_prev = (val_prev / max_share) * bar_max_w

        ax.add_patch(patches.Rectangle((bar_x, y_g1 - 0.65), w_curr, 1.05, facecolor="#ef4444", edgecolor="none"))
        ax.text(bar_x + w_curr + 0.5, y_g1 - 0.15, f"{val_curr:.1f}%", color="#b91c1c", fontsize=6.8, fontweight="bold", va="center")

        ax.add_patch(patches.Rectangle((bar_x, y_g1 - 1.95), w_prev, 0.95, facecolor="#cbd5e1", edgecolor="none"))
        ax.text(bar_x + w_prev + 0.5, y_g1 - 1.45, f"{val_prev:.1f}%", color="#64748b", fontsize=6.6, va="center")

        badge_x = 46.0
        if diff > 0:
            badge_bg = "#fee2e2"
            badge_fg = "#b91c1c"
            badge_txt = f"▲ +{diff:.1f}%"
        else:
            badge_bg = "#dcfce7"
            badge_fg = "#15803d"
            badge_txt = f"▼ {diff:.1f}%"

        ax.add_patch(patches.FancyBboxPatch((badge_x - 2.8, y_g1 - 1.6), 5.6, 2.4, boxstyle="round,pad=0.1,rounding_size=0.3", facecolor=badge_bg, edgecolor="none"))
        ax.text(badge_x, y_g1 - 0.4, badge_txt, color=badge_fg, fontsize=6.8, fontweight="bold", ha="center", va="center")
        y_g1 -= y_step_g1

    net_hrs_diff = tot_curr_mtd_hrs - tot_prev_mtd_hrs
    pct_net_diff = (net_hrs_diff / tot_prev_mtd_hrs * 100.0) if tot_prev_mtd_hrs > 0 else 0.0

    if net_hrs_diff <= 0:
        outcome_str = f"Saved {abs(net_hrs_diff):,.0f} hrs"
        tot_badge_bg = "#dcfce7"
        tot_badge_fg = "#15803d"
        tot_badge_txt = f"▼ {abs(pct_net_diff):.1f}%"
    else:
        outcome_str = f"Lost {abs(net_hrs_diff):,.0f} hrs"
        tot_badge_bg = "#fee2e2"
        tot_badge_fg = "#b91c1c"
        tot_badge_txt = f"▲ +{abs(pct_net_diff):.1f}%"

    summary_sentence = (
        f"Total downtime in first {cutoff_day} days "
        f"({curr_abbr}: {tot_curr_mtd_hrs:,.0f} vs {prev_abbr}: {tot_prev_mtd_hrs:,.0f} hrs). {outcome_str}"
    )

    ax.add_patch(patches.Rectangle((1.5, 48.0), w_box, 3.2, facecolor="#eff6ff", edgecolor="#bfdbfe", linewidth=0.6))
    ax.text(3.0, 49.6, f"Total ({curr_abbr} 1–{cutoff_day}) >>", color="#1d4ed8", fontsize=7.4, fontweight="bold", va="center")
    ax.text(17.5, 49.6, f"{tot_curr_mtd_hrs:,.1f}h", color="#1d4ed8", fontsize=7.4, fontweight="bold", ha="right", va="center")
    ax.text(19.5, 49.6, summary_sentence, color="#1e293b", fontsize=6.7, fontweight="bold", va="center")

    ax.add_patch(patches.FancyBboxPatch((46.0 - 2.8, 48.4), 5.6, 2.4, boxstyle="round,pad=0.1,rounding_size=0.3", facecolor=tot_badge_bg, edgecolor="none"))
    ax.text(46.0, 49.6, tot_badge_txt, color=tot_badge_fg, fontsize=6.8, fontweight="bold", ha="center", va="center")

    y_g2 = 82.8
    step_g2 = 2.85
    for idx, r in df_size_grid.iterrows():
        bg_c = "#f8fafc" if idx % 2 == 1 else "#ffffff"
        ax.add_patch(patches.Rectangle((50.7, y_g2 - 1.1), w_box, step_g2, facecolor=bg_c, edgecolor="none"))
        ax.plot([50.7, 98.5], [y_g2 - 1.1, y_g2 - 1.1], color="#e2e8f0", linewidth=0.45)

        ax.text(54.0, y_g2 + 0.3, str(r["MC Size"]), color="#0f172a", fontsize=7.6, fontweight="bold", va="center")
        ax.text(65.5, y_g2 + 0.3, str(r["Nos"]), color="#64748b", fontsize=7.6, va="center")
        ax.text(77.5, y_g2 + 0.3, f"{r['NPT hrs']:.1f}", color="#0f172a", fontsize=7.6, va="center")
        
        cap_val = r["NPT %"]
        t_col = "#dc2626" if cap_val >= 25 else "#0f172a"
        ax.text(90.5, y_g2 + 0.3, f"{cap_val:.1f}%", color=t_col, fontsize=7.6, fontweight="bold" if cap_val >= 25 else "normal", va="center")
        y_g2 -= step_g2

    ax.add_patch(patches.Rectangle((50.7, y_g2 - 1.1), w_box, step_g2, facecolor="#fff1f2", edgecolor="none"))
    ax.text(54.0, y_g2 + 0.3, "Summary >>", color="#e11d48", fontsize=8.0, fontweight="bold", va="center")
    ax.text(65.5, y_g2 + 0.3, f"{total_plant_mcs}", color="#0f172a", fontsize=8.0, fontweight="bold", va="center")
    ax.text(77.5, y_g2 + 0.3, f"{size_tot_hrs:,.1f}", color="#0f172a", fontsize=8.0, fontweight="bold", va="center")
    ax.text(90.5, y_g2 + 0.3, f"{size_summary_pct:.1f}%", color="#dc2626", fontsize=8.2, fontweight="bold", va="center")

    y_g3 = 38.2
    step_g3 = 4.35
    for idx, r in df_smed_grid.iterrows():
        bg_c = "#f8fafc" if idx % 2 == 1 else "#ffffff"
        ax.add_patch(patches.Rectangle((1.5, y_g3 - 2.0), w_box, step_g3, facecolor=bg_c, edgecolor="none"))
        ax.plot([1.5, 49.3], [y_g3 - 2.0, y_g3 - 2.0], color="#e2e8f0", linewidth=0.45)

        inv_wrap = "\n".join(textwrap.wrap(str(r["Involved_Mcs"]), width=32))
        ax.text(3.0, y_g3 + 0.2, str(r["Date"]), color="#0f172a", fontsize=7.4, fontweight="bold", va="center")
        ax.text(10.5, y_g3 + 0.2, str(r["Mold_Change_Qty"]), color="#0f172a", fontsize=7.4, va="center")
        ax.text(18.0, y_g3 + 0.2, f"{r['Total_Time']:.2f}", color="#0f172a", fontsize=7.4, va="center")
        ax.text(26.0, y_g3 + 0.2, f"{r['Avg_SMED']:.2f}", color="#2563eb", fontsize=7.4, fontweight="bold", va="center")
        ax.text(34.0, y_g3 + 0.2, inv_wrap, color="#475569", fontsize=6.2, va="center")
        y_g3 -= step_g3

    ax.add_patch(patches.Rectangle((1.5, y_g3 - 2.0), w_box, step_g3, facecolor="#eff6ff", edgecolor="none"))
    ax.text(3.0, y_g3 + 0.2, f"Total (1-{cutoff_day}) >>", color="#1d4ed8", fontsize=7.6, fontweight="bold", va="center")
    ax.text(10.5, y_g3 + 0.2, f"{smed_tot_qty}", color="#0f172a", fontsize=7.6, fontweight="bold", va="center")
    ax.text(18.0, y_g3 + 0.2, f"{smed_tot_time:.2f}", color="#0f172a", fontsize=7.6, fontweight="bold", va="center")
    ax.text(26.0, y_g3 + 0.2, f"{smed_avg_min:.2f}", color="#1d4ed8", fontsize=7.6, fontweight="bold", va="center")
    ax.text(34.0, y_g3 + 0.2, f"{smed_tot_qty} setups MTD (Target: 45 min)", color="#64748b", fontsize=6.8, va="center")

    y_g4 = 38.2
    step_g4 = 4.35
    for idx, r in df_maint_grid.iterrows():
        bg_c = "#fefce8" if idx % 2 == 1 else "#ffffff"
        ax.add_patch(patches.Rectangle((50.7, y_g4 - 2.0), w_box, step_g4, facecolor=bg_c, edgecolor="none"))
        ax.plot([50.7, 98.5], [y_g4 - 2.0, y_g4 - 2.0], color="#e2e8f0", linewidth=0.45)

        ax.text(x_g4[0], y_g4 + 0.2, str(r["Date"]), color="#0f172a", fontsize=7.4, fontweight="bold", va="center")
        ax.text(x_g4[1], y_g4 + 0.2, f"{r['Machine Problem*']:.1f}", color="#0284c7", fontsize=7.2, va="center")
        ax.text(x_g4[2], y_g4 + 0.2, f"{r['Robot Problem*']:.1f}", color="#0284c7", fontsize=7.2, va="center")
        ax.text(x_g4[3], y_g4 + 0.2, f"{r['Controller Problem*']:.1f}", color="#0284c7", fontsize=7.2, va="center")
        ax.text(x_g4[4], y_g4 + 0.2, f"{r['RMCS Problem*']:.1f}", color="#0284c7", fontsize=7.2, va="center")
        ax.text(x_g4[5], y_g4 + 0.2, f"{r['Oil or water Leakage*']:.1f}", color="#0284c7", fontsize=7.2, va="center")
        ax.text(x_g4[6], y_g4 + 0.2, f"{r['Total Hours']:.2f}", color="#0f172a", fontsize=7.4, fontweight="bold", va="center")
        ax.text(x_g4[7], y_g4 + 0.2, str(r["Total Share"]), color="#0f172a", fontsize=7.6, fontweight="bold", va="center")
        y_g4 -= step_g4

    ax.add_patch(patches.Rectangle((50.7, y_g4 - 2.0), w_box, step_g4, facecolor="#ecfdf5", edgecolor="none"))
    ax.text(x_g4[0], y_g4 + 0.2, maint_summary_dict["Date"], color="#047857", fontsize=7.5, fontweight="bold", va="center")
    ax.text(x_g4[1], y_g4 + 0.2, f"{maint_summary_dict['Machine Problem*']:.1f}", color="#0f172a", fontsize=7.2, fontweight="bold", va="center")
    ax.text(x_g4[2], y_g4 + 0.2, f"{maint_summary_dict['Robot Problem*']:.1f}", color="#0f172a", fontsize=7.2, fontweight="bold", va="center")
    ax.text(x_g4[3], y_g4 + 0.2, f"{maint_summary_dict['Controller Problem*']:.1f}", color="#0f172a", fontsize=7.2, fontweight="bold", va="center")
    ax.text(x_g4[4], y_g4 + 0.2, f"{maint_summary_dict['RMCS Problem*']:.1f}", color="#0f172a", fontsize=7.2, fontweight="bold", va="center")
    ax.text(x_g4[5], y_g4 + 0.2, f"{maint_summary_dict['Oil or water Leakage*']:.1f}", color="#0f172a", fontsize=7.2, fontweight="bold", va="center")
    ax.text(x_g4[6], y_g4 + 0.2, f"{maint_summary_dict['Total Hours']:.1f}", color="#047857", fontsize=7.6, fontweight="bold", va="center")
    ax.text(x_g4[7], y_g4 + 0.2, str(maint_summary_dict["Total Share"]), color="#047857", fontsize=7.8, fontweight="bold", va="center")

    plt.subplots_adjust(left=0, right=1, top=1, bottom=0)
    buf = io.BytesIO()
    plt.savefig(buf, format="jpg", facecolor=fig.get_facecolor(), edgecolor="none", dpi=240)
    plt.close(fig)
    buf.seek(0)
    return buf.getvalue()


# =========================================================
# 4. STREAMLIT RENDER CONSOLE
# =========================================================
def render_npt_module():
    c_back, c_title, c_act = st.columns([1.5, 3.5, 1.5], vertical_alignment="center")
    with c_back:
        if st.button("Back to Operations Hub", use_container_width=True):
            st.session_state["active_view"] = "hub_home"
            st.rerun()
    with c_title:
        st.markdown(
            """
            <div style="text-align:center;">
                <div style="color: #0284c7; font-size: 0.72rem; font-weight: 800; letter-spacing: 0.1em; text-transform: uppercase;">Operational Efficiency Division</div>
                <h3 style="margin:0; font-weight:800; color:#0f172a; letter-spacing: -0.01em;">NON-PRODUCTIVE TIME (NPT) ANALYTICS</h3>
            </div>
            """,
            unsafe_allow_html=True,
        )
    with c_act:
        if "m3_file_bytes" in st.session_state:
            if st.button("Change Files", use_container_width=True):
                st.session_state.pop("m3_file_bytes", None)
                st.session_state.pop("m3_sm_bytes", None)
                st.rerun()

    st.markdown("<div style='margin-bottom: 1.25rem;'></div>", unsafe_allow_html=True)

    if "m3_file_bytes" not in st.session_state:
        st.markdown(
            """
            <div style="background: #f8fafc; border: 1px solid #cbd5e1; border-left: 5px solid #0284c7; border-radius: 8px; padding: 1.2rem 1.5rem; margin-bottom: 1.5rem;">
                <div style="color: #0284c7; font-size: 0.75rem; font-weight: 800; letter-spacing: 0.08em; text-transform: uppercase; margin-bottom: 0.35rem;">
                    DATA INGESTION REQUIREMENT
                </div>
                <div style="color: #0f172a; font-size: 1.05rem; font-weight: 700; margin-bottom: 0.25rem;">
                    Operational Date Range Specification
                </div>
                <div style="color: #475569; font-size: 0.9rem; line-height: 1.5;">
                    Please upload the relevant file containing data <b>from the start of the last month to the current month's latest operational date</b>.
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        c_up1, c_up2 = st.columns(2, gap="medium")
        with c_up1:
            st.markdown(
                """
                <div style="background:#ffffff; padding:1.5rem; border-radius:8px; border:1px solid #e2e8f0; border-top:4px solid #0284c7;">
                    <h4 style="margin-top:0; color:#0f172a;">1. Primary Downtime Report (Required)</h4>
                    <p style="color:#64748b !important; font-size:0.85rem;">Ingest the multi-month ERP Downtime workbook spanning prior month to date.</p>
                </div>
                """,
                unsafe_allow_html=True,
            )
            up_dt = st.file_uploader("Select Downtime Workbook (.xlsx)", type=["xlsx", "xls"], key="up_dt_file")

        with c_up2:
            st.markdown(
                """
                <div style="background:#ffffff; padding:1.5rem; border-radius:8px; border:1px solid #e2e8f0; border-top:4px solid #10b981;">
                    <h4 style="margin-top:0; color:#0f172a;">2. Service Maintenance Report (Optional)</h4>
                    <p style="color:#64748b !important; font-size:0.85rem;">Upload workshop ticket logs for SMS token audit.</p>
                </div>
                """,
                unsafe_allow_html=True,
            )
            up_sm = st.file_uploader("Select Maintenance Ticket Workbook (.xlsx)", type=["xlsx", "xls"], key="up_sm_file")

        if up_dt is not None:
            if st.button("Ingest Workbooks & Launch Console", type="primary", use_container_width=True):
                st.session_state["m3_file_bytes"] = up_dt.getvalue()
                st.session_state["m3_sm_bytes"] = up_sm.getvalue() if up_sm is not None else None
                st.rerun()

    else:
        df_parsed = m3_parse_downtime_workbook(st.session_state["m3_file_bytes"])

        detected_sec = df_parsed["Detected_Section"].iloc[0] if "Detected_Section" in df_parsed.columns else "RIP> DPL> Plastic-3"
        sec_name, total_plant_mcs, daily_avail_hrs, pos_map, size_nos_map, unique_sizes = resolve_section_config(df_parsed, fallback_name=detected_sec)

        section_display_name = sec_name.split(">")[-1].strip()

        # Grouped Industrial Toolbar
        st.markdown('<div class="control-bar-card">', unsafe_allow_html=True)
        c_flr, c_date, c_excl, c_causes, c_snap = st.columns([1.1, 1.3, 1.8, 1.4, 1.4], gap="small", vertical_alignment="bottom")

        with c_flr:
            floor_choice = st.radio("🏢 Floor View:", ["ALL", "FF", "GF"], horizontal=True, key="m3_floor_view")

        if floor_choice == "FF":
            df_scoped = df_parsed[df_parsed["Floor"] == "FF"].copy()
            total_plant_mcs = 25
            daily_avail_hrs = 25 * 24.0
        elif floor_choice == "GF":
            df_scoped = df_parsed[df_parsed["Floor"] == "GF"].copy()
            total_plant_mcs = 36
            daily_avail_hrs = 36 * 24.0
        else:
            df_scoped = df_parsed.copy()

        raw_start_dates = (df_scoped["From_DT"].dropna() - pd.Timedelta(hours=8)).dt.normalize()
        avail_dates_list = sorted(raw_start_dates.unique().tolist())
        avail_cutoff_strs = [d.strftime("%Y-%m-%d") for d in avail_dates_list]

        with c_date:
            sel_cutoff_str = st.selectbox(
                f"Cutoff Date (8 AM–8 AM)",
                avail_cutoff_strs,
                index=len(avail_cutoff_strs) - 1 if avail_cutoff_strs else 0,
            )

        all_present_causes = sorted([c for c in df_scoped["Cause"].dropna().unique()])
        default_excluded = [c for c in all_present_causes if "server error" in c.lower()]

        with c_excl:
            excluded_causes = st.multiselect(
                "🚫 Drop Causes from Analysis:",
                all_present_causes,
                default=default_excluded,
                help="Excluded causes will be omitted from total NPT hours and capacity loss metrics.",
            )

        sel_date_obj = pd.to_datetime(sel_cutoff_str)
        cutoff_day = sel_date_obj.day
        day_formatted = sel_date_obj.strftime("%d-%b")

        cutoff_eval_dt = sel_date_obj + pd.Timedelta(days=1, hours=8)
        df_downtime_raw = slice_downtime_by_operational_days(df_scoped, cutoff_cutoff_dt=cutoff_eval_dt)

        # Apply Cause Exclusion Filter
        if excluded_causes:
            df_downtime = df_downtime_raw[~df_downtime_raw["Cause"].isin(excluded_causes)].copy()
        else:
            df_downtime = df_downtime_raw.copy()

        all_months = sorted(df_downtime["YearMonth"].dropna().unique())
        active_month = sel_date_obj.to_period("M")
        if active_month not in all_months and len(all_months) > 0:
            active_month = all_months[-1]

        active_m_df = df_downtime[df_downtime["YearMonth"] == active_month]
        curr_month_name = sel_date_obj.strftime("%B")

        valid_causes_for_picker = sorted([c for c in df_downtime["Cause"].dropna().unique() if "Server Error" not in str(c)])
        
        if "top_10_causes_selected" not in st.session_state:
            st.session_state["top_10_causes_selected"] = [c for c in DEFAULT_TOP_10_CAUSES if c in valid_causes_for_picker]

        with c_causes:
            with st.popover(f"Top Causes ({len(st.session_state['top_10_causes_selected'])}/10 Selected)"):
                st.markdown("##### Choose up to 10 Causes to Compare")
                st.caption("Selected causes are tracked in impact charts and reports:")
                
                updated_selection = []
                for c_item in valid_causes_for_picker:
                    is_currently_checked = c_item in st.session_state["top_10_causes_selected"]
                    chk = st.checkbox(c_item, value=is_currently_checked, key=f"chk_cause_{c_item}")
                    if chk:
                        updated_selection.append(c_item)

                if len(updated_selection) > 10:
                    st.warning("Maximum 10 causes allowed! Keeping first 10 selected.")
                    updated_selection = updated_selection[:10]

                if st.button("Apply Selected Causes", type="primary", use_container_width=True):
                    st.session_state["top_10_causes_selected"] = updated_selection
                    st.rerun()

        selected_top_causes = st.session_state["top_10_causes_selected"]

        df_last_day = active_m_df[active_m_df["DateStr"] == sel_cutoff_str].copy()
        df_mtd = active_m_df[active_m_df["DayNum"] <= cutoff_day].copy()

        month_idx = all_months.index(active_month) if active_month in all_months else -1
        if month_idx > 0:
            prev_month = all_months[month_idx - 1]
            prev_m_df = df_downtime[df_downtime["YearMonth"] == prev_month]
            prev_month_name = prev_m_df["MonthName"].iloc[0]
            prev_m_mtd = prev_m_df[prev_m_df["DayNum"] <= cutoff_day].copy()
        else:
            prev_month = active_month
            prev_m_df = active_m_df
            prev_month_name = "Prior"
            prev_m_mtd = df_mtd

        curr_abbr = curr_month_name[:3].capitalize()
        prev_abbr = prev_month_name[:3].capitalize()

        # Cumulative MTD Hours: Accumulates every single day
        mtd_daily_capped = df_mtd.groupby(["DateClean", "Machine"])["Hours"].sum().clip(upper=24.0).reset_index()
        tot_curr_mtd_hrs = mtd_daily_capped["Hours"].sum()

        prev_daily_capped = prev_m_mtd.groupby(["DateClean", "Machine"])["Hours"].sum().clip(upper=24.0).reset_index()
        tot_prev_mtd_hrs = prev_daily_capped["Hours"].sum()

        tot_avail_period = total_plant_mcs * 24.0 * max(1, cutoff_day)
        curr_mtd_npt_pct = (tot_curr_mtd_hrs / tot_avail_period * 100.0) if tot_avail_period > 0 else 0.0
        prev_mtd_npt_pct = (tot_prev_mtd_hrs / tot_avail_period * 100.0) if tot_avail_period > 0 else 0.0

        # Strict Single-Day Physics Cap: Maximum 24.0 Hours per Machine
        last_day_mc_hrs = df_last_day.groupby("Machine")["Hours"].sum().clip(upper=24.0)
        last_day_total_hrs = last_day_mc_hrs.sum()
        last_day_avail = daily_avail_hrs
        last_day_npt_pct = (last_day_total_hrs / last_day_avail * 100.0) if last_day_avail > 0 else 0.0

        curr_share_dict = (
            (df_mtd.groupby("Cause")["Hours"].sum() / tot_curr_mtd_hrs * 100).to_dict()
            if tot_curr_mtd_hrs > 0
            else {}
        )
        prev_share_dict = (
            (prev_m_mtd.groupby("Cause")["Hours"].sum() / tot_prev_mtd_hrs * 100).to_dict()
            if tot_prev_mtd_hrs > 0
            else {}
        )
        curr_hours_dict = df_mtd.groupby("Cause")["Hours"].sum().to_dict()

        df_size_grid, size_tot_hrs, size_summary_pct = m3_compute_size_wise_npt(df_mtd, cutoff_day, total_plant_mcs, size_nos_map, unique_sizes)
        df_smed_grid, smed_tot_qty, smed_tot_time, smed_avg_min = m3_compute_smed_table(df_mtd, max_days=7)
        df_maint_grid, maint_summary_dict = m3_compute_maint_daily_table(df_mtd, cutoff_day, daily_avail_hrs, total_plant_mcs, max_days=7)

        jpg_bytes = m3_generate_2x2_executive_jpg(
            sel_date_obj,
            cutoff_day,
            selected_top_causes,
            curr_share_dict,
            prev_share_dict,
            curr_hours_dict,
            tot_curr_mtd_hrs,
            tot_prev_mtd_hrs,
            curr_month_name,
            prev_month_name,
            df_size_grid,
            size_tot_hrs,
            size_summary_pct,
            df_smed_grid,
            smed_tot_qty,
            smed_tot_time,
            smed_avg_min,
            df_maint_grid,
            maint_summary_dict,
            section_display_name,
            total_plant_mcs,
        )

        with c_snap:
            st.download_button(
                label="Download 2×2 JPG Report",
                data=jpg_bytes,
                file_name=f"NPT_4Grid_Report_{section_display_name}_{sel_cutoff_str}.jpg",
                mime="image/jpeg",
                use_container_width=True,
            )
        st.markdown("</div>", unsafe_allow_html=True)

        # 4 Standard KPI Cards Matching Scrap Module Styling
        k1, k2, k3, k4 = st.columns(4)
        k1.markdown(
            f'<div class="kpi-card blue"><div class="kpi-title">{curr_abbr.upper()} 01–{cutoff_day:02d} TOTAL NPT</div>'
            f'<div class="kpi-val">{tot_curr_mtd_hrs:,.1f} H</div>'
            f'<div class="kpi-sub">{curr_mtd_npt_pct:.2f}% Section NPT ({tot_curr_mtd_hrs/cutoff_day:.1f} H/Day)</div></div>',
            unsafe_allow_html=True,
        )
        k2.markdown(
            f'<div class="kpi-card purple"><div class="kpi-title">{section_display_name.upper()} CAPACITY NPT %</div>'
            f'<div class="kpi-val">{size_summary_pct:.2f}%</div>'
            f'<div class="kpi-sub">Of {tot_avail_period:,.0f} H Total ({total_plant_mcs} MCs)</div></div>',
            unsafe_allow_html=True,
        )
        
        ongoing_count = df_last_day["Slice_Ongoing"].sum()
        k3.markdown(
            f'<div class="kpi-card red"><div class="kpi-title">LAST DAY NPT ({day_formatted})</div>'
            f'<div class="kpi-val">{last_day_total_hrs:.1f} H</div>'
            f'<div class="kpi-sub">{last_day_npt_pct:.2f}% Day NPT ({ongoing_count} Active)</div></div>',
            unsafe_allow_html=True,
        )
        
        maint_last_hrs = df_last_day[df_last_day["Is_Maintenance"]]["Hours"].sum()
        k4.markdown(
            f'<div class="kpi-card amber"><div class="kpi-title">LAST DAY MAINT. IMPACT</div>'
            f'<div class="kpi-val">{maint_last_hrs:.1f} H</div>'
            f'<div class="kpi-sub">{(maint_last_hrs/last_day_total_hrs*100 if last_day_total_hrs>0 else 0):.1f}% NPT Share</div></div>',
            unsafe_allow_html=True,
        )

        st.markdown("<div style='margin-bottom: 1.5rem;'></div>", unsafe_allow_html=True)

        col_left, col_right = st.columns([1.55, 0.95], gap="large")

        with col_left:
            st.markdown(
                f"""
                <div class="section-panel-header">
                    <h4 class="section-panel-title">{section_display_name.upper()} MACHINE NPT INCIDENTS LOG</h4>
                    <span style="color: #64748b; font-size: 0.8rem; font-weight: 600;">{day_formatted} (8 AM–8 AM)</span>
                </div>
                """,
                unsafe_allow_html=True,
            )
            df_cons_log = m3_compute_consolidated_daily_log(df_last_day)
            if not df_cons_log.empty:
                display_df = df_cons_log[["Position", "Machine", "Department", "Status", "Combined Causes", "Duration", "Start Time"]]
                st.dataframe(display_df, use_container_width=True, hide_index=True, height=390)
                st.download_button(
                    "📥 Export Daily Incident Log (.xlsx)",
                    convert_df_to_styled_excel(display_df, "Daily_Incidents"),
                    f"Daily_Incidents_{sel_cutoff_str}.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            else:
                st.success("Zero downtime logged for this date.")

        with col_right:
            top_mtd_causes = df_mtd.groupby("Cause")["Hours"].sum().sort_values(ascending=False).head(4)
            top_causes_lines = []
            for c_n, c_h in top_mtd_causes.items():
                c_clean = c_n.replace("*", "").strip()
                c_pct = (c_h / tot_curr_mtd_hrs * 100) if tot_curr_mtd_hrs > 0 else 0.0
                top_causes_lines.append(f"• {c_clean}: {c_h:,.2f} Hrs ({c_pct:.2f}% share)")

            maint_items = [
                "Machine Problem*",
                "Robot Problem*",
                "Controller Problem*",
                "RMCS Problem*",
                "Oil or water Leakage*",
            ]
            maint_lines = []
            tot_tech_day = 0.0
            for mi in maint_items:
                mi_clean = mi.replace("*", "").strip()
                mi_h = df_last_day[df_last_day["Cause"] == mi]["Hours"].sum()
                tot_tech_day += mi_h
                maint_lines.append(f"• {mi_clean}: {mi_h:.2f} Hrs")

            smed_last_df = df_last_day[df_last_day["Is_SMED"]]
            smed_setups_last = len(smed_last_df)
            smed_hrs_last = smed_last_df["Hours"].sum()
            smed_avg_min_last = (smed_hrs_last / smed_setups_last * 60.0) if smed_setups_last > 0 else 0.0
            smed_mcs_str = ", ".join(sorted(set(str(v) for v in smed_last_df["Position"].unique() if str(v) != "-"))) if smed_setups_last > 0 else "None"

            curr_abbr_txt = curr_month_name[:3].capitalize()
            prev_abbr_txt = prev_month_name[:3].capitalize()
            comp_like_for_like_str = f"vs. {tot_prev_mtd_hrs:,.2f} Hrs ({prev_mtd_npt_pct:.2f}% NPT) in {prev_abbr_txt} 01–{cutoff_day:02d}"

            # WhatsApp briefing preserves emojis
            whatsapp_msg = f"""📅 *Date:* {sel_date_obj.strftime('%d-%m-%Y')}

Dear Sir,

*1. Overall Monthly NPT Analysis (MTD Like-for-Like: Day 1–{cutoff_day:02d})*

Current Month Total NPT ({curr_abbr_txt} 01–{cutoff_day:02d}): *{tot_curr_mtd_hrs:,.2f} Hours ({curr_mtd_npt_pct:.2f}% NPT)* ({comp_like_for_like_str}).

*Top Contributing Causes ({curr_month_name}):*
{chr(10).join(top_causes_lines)}

*2. Machine Maintenance & Technical NPT (Last Day: {day_formatted})*
{chr(10).join(maint_lines)}
*Total Maintenance Impact:* ~{tot_tech_day:.2f} Hrs ({(tot_tech_day/last_day_total_hrs*100 if last_day_total_hrs>0 else 0):.0f}% of daily NPT | {(tot_tech_day/daily_avail_hrs*100):.2f}% daily section capacity)

*3. Mold Change & SMED Performance (Last Day: {day_formatted})*
• Mold Changes Completed: *{smed_setups_last} setups*
• Total Setup Time: *{smed_hrs_last:.2f} Hours*
• Average SMED: *{smed_avg_min_last:.2f} Min/change*
• Involved Machines: {smed_mcs_str}
• Month-to-Date SMED: *{smed_tot_qty} setups* completed totaling *{smed_tot_time:.2f} Hours* (MTD Avg: *{smed_avg_min:.2f} Min*)"""

            st.markdown(
                """
                <div class="section-panel-header">
                    <h4 class="section-panel-title">EXECUTIVE BRIEFING TEXT</h4>
                    <span style="color: #64748b; font-size: 0.8rem; font-weight: 600;">Daily Stoppage Summary</span>
                </div>
                """,
                unsafe_allow_html=True,
            )
            st.markdown(
                f"""<div class="narrative-block">
                    <p style="margin:0 0 0.5rem 0; font-weight:800; color:#1e293b;">{section_display_name.upper()} DAILY NPT & DOWNTIME BRIEF</p>
                    <p style="margin:0 0 0.75rem 0; color:#64748b; font-size:0.82rem;"><b>Date:</b> {sel_date_obj.strftime('%d-%m-%Y')}</p>
                    <h5>1. Overall Monthly NPT Analysis (MTD Like-for-Like: Day 1–{cutoff_day:02d})</h5>
                    <p>Current Month Total NPT ({curr_abbr_txt} 01–{cutoff_day:02d}): <b>{tot_curr_mtd_hrs:,.2f} Hours ({curr_mtd_npt_pct:.2f}% NPT)</b> ({comp_like_for_like_str}).</p>
                    <p style="margin:0.25rem 0 0.5rem 0;">{'<br>'.join(top_causes_lines)}</p>
                    <h5>2. Machine Maintenance & Technical NPT (Last Day: {day_formatted})</h5>
                    <p style="margin:0.25rem 0 0.5rem 0;">{'<br>'.join(maint_lines)}<br><b>Total Maintenance Impact:</b> ~{tot_tech_day:.2f} Hrs ({(tot_tech_day/daily_avail_hrs*100):.2f}% daily section capacity)</p>
                    <h5>3. Mold Change & SMED Performance (Last Day: {day_formatted})</h5>
                    <p>• Mold Changes Completed: <b>{smed_setups_last} setups</b> ({smed_hrs_last:.2f} Hours | Avg: <b>{smed_avg_min_last:.2f} Min/change</b>)<br>
                    • Involved Machines: {smed_mcs_str}<br>
                    • Month-to-Date SMED: <b>{smed_tot_qty} setups</b> completed totaling <b>{smed_tot_time:.2f} Hours</b> (MTD Avg: <b>{smed_avg_min:.2f} Min</b>)</p>
                </div>""",
                unsafe_allow_html=True,
            )

            with st.expander("Copy Plain Text Brief for WhatsApp"):
                st.text_area("Brief Text", value=whatsapp_msg, height=200, label_visibility="collapsed")

        st.divider()

        # Detailed Analytical Tabs with MoM Support & New Category Pareto
        tab_size, tab_smed, tab_maint, tab_pareto = st.tabs([
            "MC Size-Wise Capacity Loss",
            "SMED (Mold Changeover)",
            "Maintenance & SMS Audit",
            "Category Pareto & Cause Comparison",
        ])

        with tab_size:
            c_sz_title, c_sz_tgl, c_sz_chk = st.columns([2.5, 1.8, 1.4], vertical_alignment="center")
            with c_sz_title:
                st.markdown(f"#### {section_display_name.upper()} SIZE-WISE CAPACITY LOSS (DAY 1–{cutoff_day})")
            with c_sz_tgl:
                show_size_mom = st.toggle(
                    f"Compare with Prior Month ({curr_abbr} 01–{cutoff_day:02d} vs {prev_abbr} 01–{cutoff_day:02d})",
                    value=False,
                    key="tgl_size_mom",
                )
            with c_sz_chk:
                show_size_variances = False
                if show_size_mom:
                    show_size_variances = st.checkbox("Show Variance Details", value=False, key="chk_sz_var")

            if show_size_mom:
                st.caption(f"Showing Side-by-Side Size Comparison: **{curr_abbr} 01–{cutoff_day:02d}** vs **{prev_abbr} 01–{cutoff_day:02d}**")
                df_size_mom = m3_compute_size_wise_mom(
                    df_mtd, prev_m_mtd, cutoff_day, total_plant_mcs, size_nos_map, unique_sizes,
                    prev_abbr=prev_abbr, curr_abbr=curr_abbr, include_variances=show_size_variances,
                )
                st.dataframe(df_size_mom, use_container_width=True, hide_index=True)
                st.download_button(
                    "📥 Export MoM Size Comparison (.xlsx)",
                    convert_df_to_styled_excel(df_size_mom, "Size_MoM_Comparison"),
                    f"Size_MoM_Comparison_{sel_cutoff_str}.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            else:
                st.dataframe(df_size_grid, use_container_width=True, hide_index=True)
                st.download_button(
                    "📥 Export Size NPT Table (.xlsx)",
                    convert_df_to_styled_excel(df_size_grid, "Size_Wise_NPT"),
                    f"Size_Wise_NPT_{sel_cutoff_str}.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )

        with tab_smed:
            st.markdown(f"#### SMED CHANGEOVER PERFORMANCE (DAY 1–{cutoff_day})")
            st.dataframe(df_smed_grid, use_container_width=True, hide_index=True)
            st.download_button(
                "📥 Export SMED Daily Performance (.xlsx)",
                convert_df_to_styled_excel(df_smed_grid, "SMED_Daily"),
                f"SMED_Performance_{sel_cutoff_str}.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        with tab_maint:
            st.markdown(f"#### TECHNICAL & BREAKDOWN TREND (DAY 1–{cutoff_day})")
            st.dataframe(df_maint_grid, use_container_width=True, hide_index=True)
            st.download_button(
                "📥 Export Technical Maintenance Log (.xlsx)",
                convert_df_to_styled_excel(df_maint_grid, "Maintenance_Trend"),
                f"Maintenance_Trend_{sel_cutoff_str}.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        with tab_pareto:
            st.markdown(f"#### 80/20 PARETO ANALYSIS & CAUSE LEDGER (DAY 1–{cutoff_day})")
            cause_agg = df_mtd.groupby(["Category", "CauseClean"])["Hours"].sum().reset_index()
            cause_agg = cause_agg.sort_values("Hours", ascending=False).reset_index(drop=True)
            cause_agg["Share %"] = (cause_agg["Hours"] / tot_curr_mtd_hrs * 100.0) if tot_curr_mtd_hrs > 0 else 0.0
            cause_agg["Cum %"] = cause_agg["Share %"].cumsum()

            c_par1, c_par2 = st.columns([1.6, 1.4])
            with c_par1:
                fig_pareto = make_subplots(specs=[[{"secondary_y": True}]])
                fig_pareto.add_trace(
                    go.Bar(
                        x=cause_agg["CauseClean"][:12],
                        y=cause_agg["Hours"][:12],
                        name="Downtime (Hrs)",
                        marker_color="#0284c7",
                    ),
                    secondary_y=False,
                )
                fig_pareto.add_trace(
                    go.Scatter(
                        x=cause_agg["CauseClean"][:12],
                        y=cause_agg["Cum %"][:12],
                        name="Cumulative %",
                        mode="lines+markers",
                        marker_color="#ef4444",
                        line=dict(width=2),
                    ),
                    secondary_y=True,
                )
                fig_pareto.update_layout(
                    title="Top Downtime Drivers (Pareto 80/20)",
                    xaxis_tickangle=-45,
                    height=360,
                    margin=dict(l=20, r=20, t=40, b=20),
                    legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="right", x=1),
                )
                fig_pareto.update_yaxes(title_text="Hours Lost", secondary_y=False)
                fig_pareto.update_yaxes(title_text="Cumulative %", range=[0, 105], secondary_y=True)
                st.plotly_chart(fig_pareto, use_container_width=True)

            with c_par2:
                cat_agg = df_mtd.groupby("Category")["Hours"].sum().reset_index()
                cat_agg["Share %"] = (cat_agg["Hours"] / tot_curr_mtd_hrs * 100.0) if tot_curr_mtd_hrs > 0 else 0.0
                fig_cat = px.pie(
                    cat_agg,
                    names="Category",
                    values="Hours",
                    title="Downtime by Industrial Department",
                    color_discrete_sequence=px.colors.qualitative.Prism,
                    hole=0.45,
                )
                fig_cat.update_layout(height=360, margin=dict(l=20, r=20, t=40, b=20))
                st.plotly_chart(fig_cat, use_container_width=True)

            st.dataframe(cause_agg, use_container_width=True, hide_index=True)
            st.download_button(
                "📥 Export Full Cause Ledger (.xlsx)",
                convert_df_to_styled_excel(cause_agg, "Cause_Ledger"),
                f"Full_Cause_Ledger_{sel_cutoff_str}.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )


# Execute module
render_npt_module()
