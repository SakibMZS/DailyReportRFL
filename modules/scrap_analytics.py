import io
import os
import re
import sys
import textwrap
import streamlit as st
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import matplotlib.pyplot as plt
import matplotlib.patches as patches

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

from config import resolve_section_config


def clean_mold_name(val):
    if not val or pd.isna(val):
        return "-"
    text = str(val).strip()
    text = re.sub(r"\(.*?\)", "", text).strip()
    return text if text else "-"


def get_col(df, candidates, default=None):
    for c in candidates:
        if c in df.columns:
            return c
    return default


@st.cache_data
def m2_parse_workbook(file_bytes):
    file_stream = io.BytesIO(file_bytes)
    xls = pd.ExcelFile(file_stream)

    sheet_name = (
        "RejectionReport"
        if "RejectionReport" in xls.sheet_names
        else ("Rejection" if "Rejection" in xls.sheet_names else ("This Month" if "This Month" in xls.sheet_names else xls.sheet_names[0]))
    )
    df_raw = pd.read_excel(xls, sheet_name=sheet_name, header=None)

    header_idx = None
    for idx, row in df_raw.iterrows():
        row_str = " ".join([str(v) for v in row.values])
        if "Machine" in row_str and (
            "Quantity" in row_str or "Qty" in row_str or "Cause" in row_str
        ):
            header_idx = idx
            break

    df_clean = (
        pd.read_excel(xls, sheet_name=sheet_name, skiprows=header_idx)
        if header_idx is not None
        else pd.read_excel(xls, sheet_name=sheet_name)
    )
    df_clean.columns = [str(c).strip() for c in df_clean.columns]

    date_col = get_col(
        df_clean, ["Added Date", "Date", "Entry Date", "AddedDate"], df_clean.columns[-1]
    )
    df_clean["DateClean"] = pd.to_datetime(df_clean[date_col], errors="coerce")
    df_clean = df_clean.dropna(subset=["DateClean"]).sort_values("DateClean")
    df_clean["DateStr"] = df_clean["DateClean"].dt.strftime("%Y-%m-%d")
    df_clean["YearMonth"] = df_clean["DateClean"].dt.to_period("M")
    df_clean["Day"] = df_clean["DateClean"].dt.day

    entry_hours = df_clean["DateClean"].dt.hour
    df_clean["Shift"] = entry_hours.apply(lambda h: "Shift A (Day)" if 12 <= h <= 23 else "Shift B (Night)")

    qty_col = get_col(df_clean, ["Quantity", "Qty", "Rejection Pcs", "Qty (Pcs)"])
    if qty_col:
        df_clean["Qty_Pcs"] = (pd.to_numeric(df_clean[qty_col], errors="coerce").fillna(0.0) * 1000.0).round()
    else:
        df_clean["Qty_Pcs"] = 0.0

    wt_col = get_col(df_clean, ["Weight", "Rejection Ton", "Weight (Ton)", "Weight (kg)"])
    if wt_col:
        df_clean["Weight_Ton"] = pd.to_numeric(df_clean[wt_col], errors="coerce").fillna(0.0)
    else:
        df_clean["Weight_Ton"] = 0.0

    sec_name, total_mcs, daily_avail_hrs, pos_map, size_counts, unique_sizes = resolve_section_config(df_clean)
    df_clean["Detected_Section"] = sec_name

    mc_col = get_col(df_clean, ["Machine", "MC SL"], "Machine")
    df_clean["Position"] = df_clean[mc_col].astype(str).map(pos_map).fillna(df_clean[mc_col].astype(str))

    unique_months = sorted(df_clean["YearMonth"].unique())
    if len(unique_months) >= 2:
        df_prev = df_clean[df_clean["YearMonth"] == unique_months[-2]].copy()
        df_curr = df_clean[df_clean["YearMonth"] == unique_months[-1]].copy()
    else:
        df_prev = pd.DataFrame()
        df_curr = df_clean.copy()

    return df_prev, df_curr, df_clean


def m2_compute_daily_rejection(df_day, pos_map, min_qty=50):
    if df_day.empty:
        return pd.DataFrame()

    mc_col = get_col(df_day, ["Machine", "MC SL", "MC Name"], df_day.columns[2])
    item_col = get_col(df_day, ["Item", "Mold", "Item Name", "Mold / Item"], df_day.columns[3])
    cause_col = get_col(df_day, ["Cause", "Causes", "Defect", "Reason"], "Cause")

    records = []
    for mc, grp in df_day.groupby(mc_col):
        total_pcs = grp["Qty_Pcs"].sum()
        total_ton = grp["Weight_Ton"].sum()

        if cause_col in grp.columns:
            causes_list = [
                str(c).strip().replace("*", "")
                for c in grp[cause_col].dropna().unique()
                if str(c).strip()
            ]
            causes_str = ", ".join(causes_list) if causes_list else "No Rejection"
        else:
            causes_str = "-"

        raw_mold = (
            str(grp[item_col].iloc[0])
            if (item_col in grp.columns and not grp[item_col].dropna().empty)
            else "-"
        )
        mold_name = clean_mold_name(raw_mold)
        pos = pos_map.get(str(mc), str(mc))

        if total_pcs >= min_qty:
            records.append({
                "Position": pos,
                "Machine": str(mc),
                "Causes": causes_str,
                "Qty": int(round(total_pcs)),
                "Weight (Ton)": round(total_ton, 4),
                "Weight (kg)": round(total_ton * 1000.0, 2),
                "Mold": mold_name,
            })

    df_res = pd.DataFrame(records)
    if not df_res.empty:
        df_res = df_res.sort_values("Qty", ascending=False).reset_index(drop=True)
    return df_res


def m2_export_rejection_excel(df_day_filtered):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Critical Rejection Log"
    ws.views.sheetView[0].showGridLines = True

    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="000000")
    data_font = Font(name="Calibri", size=10, color="000000")
    summary_font = Font(name="Calibri", size=11, bold=True, color="DC2626")
    summary_qty_font = Font(name="Calibri", size=11, bold=True, color="000000")

    thin_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )

    headers = ["Position", "Machine", "Causes", "Qty", "Mold"]
    ws.append(headers)

    for col_num in range(1, 6):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = yellow_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_num == 4 else "left", vertical="center")
        cell.border = thin_border

    ws.row_dimensions[1].height = 24

    current_row = 2
    for _, row in df_day_filtered.iterrows():
        ws.cell(row=current_row, column=1, value=str(row["Position"])).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=current_row, column=2, value=str(row["Machine"])).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=current_row, column=3, value=str(row["Causes"])).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=current_row, column=4, value=int(row["Qty"])).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=current_row, column=5, value=str(row["Mold"])).alignment = Alignment(horizontal="left", vertical="center")

        for col_num in range(1, 6):
            c = ws.cell(row=current_row, column=col_num)
            c.font = data_font
            c.border = thin_border

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    ws.cell(row=current_row, column=1, value="")
    ws.cell(row=current_row, column=2, value="")
    sum_cell = ws.cell(row=current_row, column=3, value="Summary >>")
    sum_cell.font = summary_font
    sum_cell.alignment = Alignment(horizontal="center", vertical="center")

    qty_sum_cell = ws.cell(row=current_row, column=4, value=f"=SUM(D2:D{current_row-1})")
    qty_sum_cell.font = summary_qty_font
    qty_sum_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.cell(row=current_row, column=5, value="")

    for col_num in range(1, 6):
        ws.cell(row=current_row, column=col_num).border = thin_border

    ws.row_dimensions[current_row].height = 22

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def m2_compute_cause_breakdown(df_scope):
    if df_scope.empty:
        return pd.DataFrame()
    cause_col = get_col(df_scope, ["Cause", "Causes", "Defect", "Reason"], "Cause")
    mc_col = get_col(df_scope, ["Machine", "MC SL"], "Machine")

    res = (
        df_scope.groupby(cause_col)
        .agg(
            Rej_Pcs=("Qty_Pcs", "sum"),
            Rej_Ton=("Weight_Ton", "sum"),
            Entries_Count=(cause_col, "count"),
            MC_Count=(mc_col, "nunique") if mc_col in df_scope.columns else (cause_col, "count"),
        )
        .reset_index()
    )

    res["Rej_Pcs"] = res["Rej_Pcs"].round().astype(int)
    res["Rej_Kg"] = (res["Rej_Ton"] * 1000.0).round(1)
    res["Rej_Ton"] = res["Rej_Ton"].round(4)
    res["Cause"] = res[cause_col].astype(str).str.replace("*", "", regex=False).str.strip()

    tot_pcs = res["Rej_Pcs"].sum()
    res["% Share Raw"] = (res["Rej_Pcs"] / tot_pcs * 100.0).round(1) if tot_pcs > 0 else 0.0
    res = res.sort_values("Rej_Pcs", ascending=False).reset_index(drop=True)
    res["% Share"] = res["% Share Raw"].apply(lambda x: f"{x:.1f}%")
    return res


def m2_compute_cause_mom_comparison(df_prev_scope, df_curr_scope, prev_abbr, curr_abbr, include_variances=False):
    """Computes like-for-like MoM variance at the defect cause level with paired columns."""
    cause_col = "Cause"
    grp_c = df_curr_scope.groupby(cause_col).agg(
        Curr_Pcs=("Qty_Pcs", "sum"),
        Curr_Ton=("Weight_Ton", "sum"),
        Curr_Entries=("Qty_Pcs", "count"),
    )
    grp_p = (
        df_prev_scope.groupby(cause_col).agg(
            Prev_Pcs=("Qty_Pcs", "sum"),
            Prev_Ton=("Weight_Ton", "sum"),
            Prev_Entries=("Qty_Pcs", "count"),
        )
        if not df_prev_scope.empty
        else pd.DataFrame(columns=["Prev_Pcs", "Prev_Ton", "Prev_Entries"])
    )

    merged = pd.merge(grp_c, grp_p, on=cause_col, how="outer").fillna(0.0)
    merged["Diff_Pcs"] = merged["Curr_Pcs"] - merged["Prev_Pcs"]
    merged["Diff_Ton"] = merged["Curr_Ton"] - merged["Prev_Ton"]
    merged["Trend_%"] = np.where(
        merged["Prev_Pcs"] > 0,
        ((merged["Diff_Pcs"] / merged["Prev_Pcs"]) * 100.0).round(1),
        0.0,
    )

    merged = merged.sort_values("Curr_Pcs", ascending=False).reset_index()
    merged["Cause"] = (
        merged["Cause"]
        .astype(str)
        .str.replace("*", "", regex=False)
        .str.strip()
    )

    c_curr_pcs = f"{curr_abbr} Pcs"
    c_prev_pcs = f"{prev_abbr} Pcs"
    c_curr_ton = f"{curr_abbr} Ton"
    c_prev_ton = f"{prev_abbr} Ton"
    c_curr_ent = f"{curr_abbr} Entries"
    c_prev_ent = f"{prev_abbr} Entries"

    merged = merged.rename(
        columns={
            "Curr_Pcs": c_curr_pcs,
            "Prev_Pcs": c_prev_pcs,
            "Curr_Ton": c_curr_ton,
            "Prev_Ton": c_prev_ton,
            "Diff_Pcs": "Variance (Pcs)",
            "Diff_Ton": "Variance (Ton)",
            "Trend_%": "MoM Trend %",
            "Curr_Entries": c_curr_ent,
            "Prev_Entries": c_prev_ent,
        }
    )

    merged[c_curr_pcs] = merged[c_curr_pcs].round().astype(int)
    merged[c_prev_pcs] = merged[c_prev_pcs].round().astype(int)
    merged["Variance (Pcs)"] = merged["Variance (Pcs)"].round().astype(int)
    merged[c_curr_ton] = merged[c_curr_ton].round(3)
    merged[c_prev_ton] = merged[c_prev_ton].round(3)
    merged["Variance (Ton)"] = merged["Variance (Ton)"].round(3)
    merged[c_curr_ent] = merged[c_curr_ent].astype(int)
    merged[c_prev_ent] = merged[c_prev_ent].astype(int)

    base_cols = [
        "Cause",
        c_curr_pcs,
        c_prev_pcs,
        c_curr_ton,
        c_prev_ton,
        c_curr_ent,
        c_prev_ent,
    ]
    if include_variances:
        base_cols = [
            "Cause",
            c_curr_pcs,
            c_prev_pcs,
            c_curr_ton,
            c_prev_ton,
            "Variance (Pcs)",
            "Variance (Ton)",
            "MoM Trend %",
            c_curr_ent,
            c_prev_ent,
        ]
    return merged[base_cols]


def m2_compute_lineman_breakdown(df_scope):
    if df_scope.empty:
        return pd.DataFrame()
    lineman_col = get_col(df_scope, ["Added By", "AddedBy", "Lineman", "Operator", "Added_By"], None)
    if not lineman_col or lineman_col not in df_scope.columns:
        return pd.DataFrame()

    mc_col = get_col(df_scope, ["Machine", "MC SL"], "Machine")

    res = (
        df_scope.groupby(lineman_col)
        .agg(
            Rej_Pcs=("Qty_Pcs", "sum"),
            Rej_Ton=("Weight_Ton", "sum"),
            Logged_Entries=(lineman_col, "count"),
            Machines_Covered=(mc_col, "nunique") if mc_col in df_scope.columns else (lineman_col, "count"),
        )
        .reset_index()
    )

    res["Rej_Pcs"] = res["Rej_Pcs"].round().astype(int)
    res["Rej_Kg"] = (res["Rej_Ton"] * 1000.0).round(1)
    res["Rej_Ton"] = res["Rej_Ton"].round(4)

    tot_pcs = res["Rej_Pcs"].sum()
    tot_ton = res["Rej_Ton"].sum()
    res["% Pcs Share Raw"] = (res["Rej_Pcs"] / tot_pcs * 100.0).round(1) if tot_pcs > 0 else 0.0
    res["% Ton Share Raw"] = (res["Rej_Ton"] / tot_ton * 100.0).round(1) if tot_ton > 0 else 0.0
    res["% Pcs Share"] = res["% Pcs Share Raw"].apply(lambda x: f"{x:.1f}%")
    res["% Ton Share"] = res["% Ton Share Raw"].apply(lambda x: f"{x:.1f}%")
    res = res.sort_values("Rej_Pcs", ascending=False).reset_index(drop=True)
    res = res.rename(columns={lineman_col: "Lineman (Added By)"})
    return res


def m2_compute_lineman_mom_comparison(
    df_prev_scope, df_curr_scope, prev_abbr, curr_abbr, selected_ops, include_variances=False
):
    """Computes like-for-like MoM variance for operators with paired columns."""
    lineman_col = "Added By"
    df_p_filt = (
        df_prev_scope[df_prev_scope[lineman_col].isin(selected_ops)]
        if not df_prev_scope.empty
        else pd.DataFrame()
    )
    df_c_filt = df_curr_scope[df_curr_scope[lineman_col].isin(selected_ops)]

    grp_c = df_c_filt.groupby(lineman_col).agg(
        Curr_Pcs=("Qty_Pcs", "sum"),
        Curr_Ton=("Weight_Ton", "sum"),
        Curr_Entries=("Qty_Pcs", "count"),
    )
    grp_p = (
        df_p_filt.groupby(lineman_col).agg(
            Prev_Pcs=("Qty_Pcs", "sum"),
            Prev_Ton=("Weight_Ton", "sum"),
            Prev_Entries=("Qty_Pcs", "count"),
        )
        if not df_p_filt.empty
        else pd.DataFrame(columns=["Prev_Pcs", "Prev_Ton", "Prev_Entries"])
    )

    merged = pd.merge(grp_c, grp_p, on=lineman_col, how="outer").fillna(0.0)
    merged["Diff_Pcs"] = merged["Curr_Pcs"] - merged["Prev_Pcs"]
    merged["Diff_Ton"] = merged["Curr_Ton"] - merged["Prev_Ton"]
    merged["Trend_%"] = np.where(
        merged["Prev_Pcs"] > 0,
        ((merged["Diff_Pcs"] / merged["Prev_Pcs"]) * 100.0).round(1),
        0.0,
    )

    merged = merged.sort_values("Curr_Pcs", ascending=False).reset_index()
    c_curr_pcs = f"{curr_abbr} Pcs"
    c_prev_pcs = f"{prev_abbr} Pcs"
    c_curr_ton = f"{curr_abbr} Ton"
    c_prev_ton = f"{prev_abbr} Ton"
    c_curr_ent = f"{curr_abbr} Entries"
    c_prev_ent = f"{prev_abbr} Entries"

    merged = merged.rename(
        columns={
            lineman_col: "Lineman (Added By)",
            "Curr_Pcs": c_curr_pcs,
            "Prev_Pcs": c_prev_pcs,
            "Curr_Ton": c_curr_ton,
            "Prev_Ton": c_prev_ton,
            "Diff_Pcs": "Variance (Pcs)",
            "Diff_Ton": "Variance (Ton)",
            "Trend_%": "MoM Trend %",
            "Curr_Entries": c_curr_ent,
            "Prev_Entries": c_prev_ent,
        }
    )

    merged[c_curr_pcs] = merged[c_curr_pcs].round().astype(int)
    merged[c_prev_pcs] = merged[c_prev_pcs].round().astype(int)
    merged["Variance (Pcs)"] = merged["Variance (Pcs)"].round().astype(int)
    merged[c_curr_ton] = merged[c_curr_ton].round(3)
    merged[c_prev_ton] = merged[c_prev_ton].round(3)
    merged["Variance (Ton)"] = merged["Variance (Ton)"].round(3)
    merged[c_curr_ent] = merged[c_curr_ent].astype(int)
    merged[c_prev_ent] = merged[c_prev_ent].astype(int)

    base_cols = [
        "Lineman (Added By)",
        c_curr_pcs,
        c_prev_pcs,
        c_curr_ton,
        c_prev_ton,
        c_curr_ent,
        c_prev_ent,
    ]
    if include_variances:
        base_cols = [
            "Lineman (Added By)",
            c_curr_pcs,
            c_prev_pcs,
            c_curr_ton,
            c_prev_ton,
            "Variance (Pcs)",
            "Variance (Ton)",
            "MoM Trend %",
            c_curr_ent,
            c_prev_ent,
        ]
    return merged[base_cols]


def m2_compute_shift_mom_comparison(df_prev_scope, df_curr_scope, prev_abbr, curr_abbr, include_variances=False):
    """Computes Shift A vs Shift B like-for-like MoM variance."""
    if df_curr_scope.empty or "Shift" not in df_curr_scope.columns:
        return pd.DataFrame()

    grp_c = df_curr_scope.groupby("Shift").agg(
        Curr_Pcs=("Qty_Pcs", "sum"),
        Curr_Ton=("Weight_Ton", "sum"),
        Curr_Entries=("Qty_Pcs", "count"),
    )
    grp_p = (
        df_prev_scope.groupby("Shift").agg(
            Prev_Pcs=("Qty_Pcs", "sum"),
            Prev_Ton=("Weight_Ton", "sum"),
            Prev_Entries=("Qty_Pcs", "count"),
        )
        if not df_prev_scope.empty
        else pd.DataFrame(columns=["Prev_Pcs", "Prev_Ton", "Prev_Entries"])
    )

    merged = pd.merge(grp_c, grp_p, on="Shift", how="outer").fillna(0.0)
    merged["Diff_Pcs"] = merged["Curr_Pcs"] - merged["Prev_Pcs"]
    merged["Diff_Ton"] = merged["Curr_Ton"] - merged["Prev_Ton"]
    merged["Trend_%"] = np.where(
        merged["Prev_Pcs"] > 0,
        ((merged["Diff_Pcs"] / merged["Prev_Pcs"]) * 100.0).round(1),
        0.0,
    )

    c_curr_pcs = f"{curr_abbr} Pcs"
    c_prev_pcs = f"{prev_abbr} Pcs"
    c_curr_ton = f"{curr_abbr} Ton"
    c_prev_ton = f"{prev_abbr} Ton"
    c_curr_ent = f"{curr_abbr} Entries"
    c_prev_ent = f"{prev_abbr} Entries"

    merged = merged.reset_index().rename(
        columns={
            "Curr_Pcs": c_curr_pcs,
            "Prev_Pcs": c_prev_pcs,
            "Curr_Ton": c_curr_ton,
            "Prev_Ton": c_prev_ton,
            "Curr_Entries": c_curr_ent,
            "Prev_Entries": c_prev_ent,
            "Diff_Pcs": "Variance (Pcs)",
            "Diff_Ton": "Variance (Ton)",
            "Trend_%": "MoM Trend %",
        }
    )

    merged[c_curr_pcs] = merged[c_curr_pcs].round().astype(int)
    merged[c_prev_pcs] = merged[c_prev_pcs].round().astype(int)
    merged["Variance (Pcs)"] = merged["Variance (Pcs)"].round().astype(int)
    merged[c_curr_ton] = merged[c_curr_ton].round(3)
    merged[c_prev_ton] = merged[c_prev_ton].round(3)
    merged["Variance (Ton)"] = merged["Variance (Ton)"].round(3)
    merged[c_curr_ent] = merged[c_curr_ent].astype(int)
    merged[c_prev_ent] = merged[c_prev_ent].astype(int)

    cols = ["Shift", c_curr_pcs, c_prev_pcs, c_curr_ton, c_prev_ton, c_curr_ent, c_prev_ent]
    if include_variances:
        cols = ["Shift", c_curr_pcs, c_prev_pcs, c_curr_ton, c_prev_ton, "Variance (Pcs)", "Variance (Ton)", "MoM Trend %", c_curr_ent, c_prev_ent]
    return merged[cols]


def m2_compute_operator_defect_pivot(df_scope):
    if df_scope.empty:
        return pd.DataFrame()
    cause_col = get_col(df_scope, ["Cause", "Causes", "Defect"], "Cause")
    lineman_col = get_col(df_scope, ["Added By", "AddedBy", "Lineman"], "Added By")

    pivot = df_scope.pivot_table(
        index=cause_col,
        columns=lineman_col,
        values="Qty_Pcs",
        aggfunc="count",
        fill_value=0,
    )
    pivot["Total Entries"] = pivot.sum(axis=1)
    pivot = pivot.sort_values("Total Entries", ascending=False).reset_index()
    pivot[cause_col] = pivot[cause_col].astype(str).str.replace("*", "", regex=False).str.strip()
    return pivot.rename(columns={cause_col: "Defect Cause Mode"})


def m2_compute_operator_rankings(df_scope):
    if df_scope.empty:
        return pd.DataFrame()
    cause_col = get_col(df_scope, ["Cause", "Causes", "Defect"], "Cause")
    lineman_col = get_col(df_scope, ["Added By", "AddedBy", "Lineman"], "Added By")

    rankings = []
    for op, grp in df_scope.groupby(lineman_col):
        cause_agg = grp.groupby(cause_col).agg(
            entry_count=("Qty_Pcs", "count"),
            rej_pcs=("Qty_Pcs", "sum"),
        ).sort_values("rej_pcs", ascending=False).head(5)

        top_items = [
            f"{c.replace('*', '').strip()} ({int(r['entry_count'])} entries, {int(round(r['rej_pcs'])):,} pcs)"
            for c, r in cause_agg.iterrows()
        ]
        row = {"Operator": op}
        for i, item in enumerate(top_items):
            row[f"Rank {i+1} Driver"] = item
        rankings.append(row)

    df_rank = pd.DataFrame(rankings)
    return df_rank


def m2_compute_daily_trend_comparison(df_prev, df_curr, prev_abbr="Aug", curr_abbr="Sep", shift_filter="All (24 Hrs)", include_variances=False):
    """
    Computes daily Day-of-Month comparison pairing pieces before tons:
    [Day of Month, Curr Pcs, Prev Pcs, Curr Ton, Prev Ton, (Optional Variances)]
    Supports 3-way shift filtering: All, Day Shift (A), Night Shift (B).
    """
    p = df_prev.copy()
    c = df_curr.copy()

    if shift_filter == "Day Shift (Shift A)":
        p = p[p["Shift"].str.contains("Shift A", na=False)] if not p.empty else p
        c = c[c["Shift"].str.contains("Shift A", na=False)] if not c.empty else c
    elif shift_filter == "Night Shift (Shift B)":
        p = p[p["Shift"].str.contains("Shift B", na=False)] if not p.empty else p
        c = c[c["Shift"].str.contains("Shift B", na=False)] if not c.empty else c

    t_prev = (
        p.groupby("Day").agg(
            Prev_Pcs=("Qty_Pcs", "sum"),
            Prev_Ton=("Weight_Ton", "sum")
        ).reset_index()
        if not p.empty
        else pd.DataFrame(columns=["Day", "Prev_Pcs", "Prev_Ton"])
    )

    t_curr = (
        c.groupby("Day").agg(
            Curr_Pcs=("Qty_Pcs", "sum"),
            Curr_Ton=("Weight_Ton", "sum")
        ).reset_index()
        if not c.empty
        else pd.DataFrame(columns=["Day", "Curr_Pcs", "Curr_Ton"])
    )

    merged = pd.merge(t_curr, t_prev, on="Day", how="outer").sort_values("Day").fillna(0.0)

    merged["Diff_Pcs"] = merged["Curr_Pcs"] - merged["Prev_Pcs"]
    merged["Diff_Ton"] = merged["Curr_Ton"] - merged["Prev_Ton"]
    merged["Trend_%"] = np.where(
        merged["Prev_Pcs"] > 0,
        ((merged["Diff_Pcs"] / merged["Prev_Pcs"]) * 100.0).round(1),
        0.0,
    )

    c_curr_pcs = f"{curr_abbr} Pcs"
    c_prev_pcs = f"{prev_abbr} Pcs"
    c_curr_ton = f"{curr_abbr} Ton"
    c_prev_ton = f"{prev_abbr} Ton"

    merged = merged.rename(
        columns={
            "Day": "Day of Month",
            "Curr_Pcs": c_curr_pcs,
            "Prev_Pcs": c_prev_pcs,
            "Curr_Ton": c_curr_ton,
            "Prev_Ton": c_prev_ton,
            "Diff_Pcs": "Variance (Pcs)",
            "Diff_Ton": "Variance (Ton)",
            "Trend_%": "MoM Trend %",
        }
    )

    merged[c_curr_pcs] = merged[c_curr_pcs].round().astype(int)
    merged[c_prev_pcs] = merged[c_prev_pcs].round().astype(int)
    merged["Variance (Pcs)"] = merged["Variance (Pcs)"].round().astype(int)
    merged[c_curr_ton] = merged[c_curr_ton].round(3)
    merged[c_prev_ton] = merged[c_prev_ton].round(3)
    merged["Variance (Ton)"] = merged["Variance (Ton)"].round(3)

    cols = ["Day of Month", c_curr_pcs, c_prev_pcs, c_curr_ton, c_prev_ton]
    if include_variances:
        cols = [
            "Day of Month",
            c_curr_pcs,
            c_prev_pcs,
            c_curr_ton,
            c_prev_ton,
            "Variance (Pcs)",
            "Variance (Ton)",
            "MoM Trend %",
        ]
    return merged[cols]


def m2_generate_scrap_jpg(
    df_day_filtered,
    sel_date_obj,
    total_rej_pcs,
    total_rej_ton,
    prev_as_of_total_ton,
    prev_as_of_avg_ton,
    curr_as_of_total_ton,
    curr_as_of_avg_ton,
    high_rej_count,
    total_day_mcs,
    top_cause,
    top_cause_pcs,
    top_cause_pct,
    top_wt_mc,
    top_wt_kg,
    top3_pct,
    top3_summary_list,
    section_label,
    total_plant_mcs,
    prev_abbr="Aug",
    curr_abbr="Sep",
    sel_day_num=1,
    min_cutoff=50,
):
    fig, ax = plt.subplots(figsize=(18, 10.5), dpi=220)
    fig.patch.set_facecolor("#f1f5f9")
    ax.set_facecolor("#f1f5f9")
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis("off")

    date_formatted = sel_date_obj.strftime("%B %d, %Y")
    day_formatted = sel_date_obj.strftime("%B %d")

    ax.text(1.5, 98.4, "DAILY REJECTION & DEFECT ANALYTICS REPORT", color="#0f172a", fontsize=16.0, fontweight="bold", va="top")
    ax.text(1.5, 95.8, f"{section_label} Machine Rejection Log (>{min_cutoff} Pcs) & Section Summary  |  Report Date: {date_formatted}", color="#64748b", fontsize=8.8, va="top")
    ax.text(98.5, 97.2, f"{section_label.upper()} OPERATIONS", color="#2563eb", fontsize=8.8, fontweight="bold", ha="right", va="top")

    kpis = [
        (f"{prev_abbr.upper()} 01–{sel_day_num:02d} TOTAL", f"{prev_as_of_total_ton:.2f} T", "Prior MTD Total", "#64748b"),
        (f"{prev_abbr.upper()} 01–{sel_day_num:02d} AVG", f"{prev_as_of_avg_ton:.2f} T/Day", "Prior Daily Baseline", "#64748b"),
        (f"{curr_abbr.upper()} 01–{sel_day_num:02d} TOTAL", f"{curr_as_of_total_ton:.2f} T", "Current MTD Total", "#2563eb"),
        (f"{curr_abbr.upper()} 01–{sel_day_num:02d} AVG", f"{curr_as_of_avg_ton:.2f} T/Day", "Current MTD Pace", "#2563eb"),
        ("LAST DAY REJECTION", f"{total_rej_ton:.3f} T", f"{total_rej_pcs:,} Pcs Lost", "#dc2626"),
        (f"CRITICAL MC (>{min_cutoff})", f"{high_rej_count} MCs", "Lines Exceeding Limit", "#7c3aed"),
    ]
    kpi_w, kpi_gap = 15.1, 1.25
    for i, (title, val, sub, col_bar) in enumerate(kpis):
        x0 = 1.5 + i * (kpi_w + kpi_gap)
        card = patches.FancyBboxPatch((x0, 87.0), kpi_w, 7.2, boxstyle="round,pad=0.15,rounding_size=0.5", facecolor="#ffffff", edgecolor="#cbd5e1", linewidth=0.8)
        ax.add_patch(card)
        top_bar = patches.FancyBboxPatch((x0 + 0.1, 93.75), kpi_w - 0.2, 0.45, boxstyle="round,pad=0.03,rounding_size=0.2", facecolor=col_bar, edgecolor="none")
        ax.add_patch(top_bar)
        ax.text(x0 + kpi_w / 2, 92.4, title, color="#64748b", fontsize=7.4, fontweight="bold", ha="center")
        ax.text(x0 + kpi_w / 2, 89.6, val, color="#0f172a", fontsize=12.5, fontweight="bold", ha="center")
        ax.text(x0 + kpi_w / 2, 87.8, sub, color="#94a3b8", fontsize=6.8, ha="center")

    left_card = patches.FancyBboxPatch((1.5, 1.5), 74.0, 84.0, boxstyle="round,pad=0.25,rounding_size=0.8", facecolor="#ffffff", edgecolor="#cbd5e1", linewidth=1)
    ax.add_patch(left_card)
    ax.text(3.5, 83.5, f"{section_label.upper()} MACHINE REJECTION LOG (>{min_cutoff} Pcs) — {day_formatted}", color="#0f172a", fontsize=11.0, fontweight="bold")
    ax.text(73.5, 83.5, f"{high_rej_count} Machines Active Above Threshold", color="#64748b", fontsize=8.0, ha="right")

    right_card = patches.FancyBboxPatch((76.5, 1.5), 22.0, 84.0, boxstyle="round,pad=0.25,rounding_size=0.8", facecolor="#ffffff", edgecolor="#cbd5e1", linewidth=1)
    ax.add_patch(right_card)
    ax.text(78.0, 83.5, "EXECUTIVE ANALYSIS", color="#0f172a", fontsize=11.0, fontweight="bold")

    n_count = len(df_day_filtered)

    if n_count <= 30:
        left_x = 2.6
        tbl_w = 71.8
        tbl_hdr = patches.Rectangle((left_x, 79.5), tbl_w, 2.6, facecolor="#1e293b", edgecolor="none")
        ax.add_patch(tbl_hdr)
        ax.text(left_x + 1.0, 80.8, "POSITION", color="#ffffff", fontsize=7.2, fontweight="bold", va="center")
        ax.text(left_x + 8.5, 80.8, "MACHINE", color="#ffffff", fontsize=7.2, fontweight="bold", va="center")
        ax.text(left_x + 18.0, 80.8, "DEFECT CAUSES", color="#ffffff", fontsize=7.2, fontweight="bold", va="center")
        ax.text(left_x + 44.0, 80.8, "QTY", color="#ffffff", fontsize=7.2, fontweight="bold", ha="right", va="center")
        ax.text(left_x + 46.0, 80.8, "MOLD / ITEM", color="#ffffff", fontsize=7.2, fontweight="bold", va="center")

        row_y = 77.2
        row_step = min(3.8, 74.0 / max(1, n_count))
        for r_i, (_, r) in enumerate(df_day_filtered.iterrows()):
            bg_c = "#f8fafc" if r_i % 2 == 1 else "#ffffff"
            row_bg = patches.Rectangle((left_x, row_y - 1.2), tbl_w, row_step, facecolor=bg_c, edgecolor="none")
            ax.add_patch(row_bg)
            ax.plot([left_x, left_x + tbl_w], [row_y - 1.2, row_y - 1.2], color="#e2e8f0", linewidth=0.45)

            ax.text(left_x + 1.0, row_y + 0.35, str(r["Position"]), color="#0f172a", fontsize=6.8, fontweight="bold", va="center")
            ax.text(left_x + 8.5, row_y + 0.35, str(r["Machine"]), color="#64748b", fontsize=6.6, va="center")
            cause_wrap = "\n".join(textwrap.wrap(str(r["Causes"]), width=38))
            ax.text(left_x + 18.0, row_y + 0.35, cause_wrap, color="#b91c1c", fontsize=6.4, va="center")
            ax.text(left_x + 44.0, row_y + 0.35, f"{int(r['Qty']):,}", color="#0f172a", fontsize=7.0, fontweight="bold", ha="right", va="center")
            mold_wrap = "\n".join(textwrap.wrap(str(r["Mold"]), width=38))
            ax.text(left_x + 46.0, row_y + 0.35, mold_wrap, color="#334155", fontsize=6.4, va="center")
            row_y -= row_step
    else:
        mid_idx = (n_count + 1) // 2
        sub_a = df_day_filtered.iloc[:mid_idx]
        sub_b = df_day_filtered.iloc[mid_idx:]
        sub_configs = [(sub_a, 2.6, 35.8), (sub_b, 39.0, 35.8)]

        for sub_df, left_x, tbl_w in sub_configs:
            tbl_hdr = patches.Rectangle((left_x, 79.5), tbl_w, 2.6, facecolor="#1e293b", edgecolor="none")
            ax.add_patch(tbl_hdr)
            ax.text(left_x + 0.8, 80.8, "POS", color="#ffffff", fontsize=7.0, fontweight="bold", va="center")
            ax.text(left_x + 4.8, 80.8, "MACHINE", color="#ffffff", fontsize=7.0, fontweight="bold", va="center")
            ax.text(left_x + 10.4, 80.8, "DEFECT CAUSES", color="#ffffff", fontsize=7.0, fontweight="bold", va="center")
            ax.text(left_x + 23.0, 80.8, "QTY", color="#ffffff", fontsize=7.0, fontweight="bold", ha="right", va="center")
            ax.text(left_x + 24.0, 80.8, "MOLD / ITEM", color="#ffffff", fontsize=7.0, fontweight="bold", va="center")

            row_y = 77.2
            row_step = 3.75
            for r_i, (_, r) in enumerate(sub_df.iterrows()):
                bg_c = "#f8fafc" if r_i % 2 == 1 else "#ffffff"
                row_bg = patches.Rectangle((left_x, row_y - 1.2), tbl_w, row_step, facecolor=bg_c, edgecolor="none")
                ax.add_patch(row_bg)
                ax.plot([left_x, left_x + tbl_w], [row_y - 1.2, row_y - 1.2], color="#e2e8f0", linewidth=0.45)

                ax.text(left_x + 0.8, row_y + 0.35, str(r["Position"]), color="#0f172a", fontsize=6.7, fontweight="bold", va="center")
                ax.text(left_x + 4.8, row_y + 0.35, str(r["Machine"]), color="#64748b", fontsize=6.5, va="center")
                cause_wrap = "\n".join(textwrap.wrap(str(r["Causes"]), width=20))
                ax.text(left_x + 10.4, row_y + 0.35, cause_wrap, color="#b91c1c", fontsize=6.2, va="center")
                ax.text(left_x + 23.0, row_y + 0.35, f"{int(r['Qty']):,}", color="#0f172a", fontsize=6.8, fontweight="bold", ha="right", va="center")
                mold_wrap = "\n".join(textwrap.wrap(str(r["Mold"]), width=20))
                ax.text(left_x + 24.0, row_y + 0.35, mold_wrap, color="#334155", fontsize=6.2, va="center")
                row_y -= row_step

    c1 = patches.FancyBboxPatch((77.5, 42.5), 20.0, 39.0, boxstyle="round,pad=0.2,rounding_size=0.5", facecolor="#fff7f7", edgecolor="#fecaca", linewidth=0.8)
    ax.add_patch(c1)
    ax.text(78.6, 78.8, "Rejection Pareto & Top Causes", color="#b91c1c", fontsize=9.6, fontweight="bold")

    top3_lines = "\n".join([f"  {idx+1}. {c}: {p:,} pcs ({pct:.1f}%)" for idx, (c, p, pct) in enumerate(top3_summary_list)])
    t1 = (
        f"• Top 3 Causes ({top3_pct:.1f}% of loss):\n"
        f"{top3_lines}\n\n"
        f"• Heaviest Loss Machine:\n"
        f"  {top_wt_mc} ({top_wt_kg:.1f} kg loss).\n\n"
        f"• Critical Observations:\n"
        f"  Repetitive short filling noted\n"
        f"  on active production molds.\n"
        f"  Color change purging requires\n"
        f"  strict standardization."
    )
    ax.text(78.6, 75.2, t1, color="#7f1d1d", fontsize=8.2, linespacing=1.45, va="top")

    c2 = patches.FancyBboxPatch((77.5, 2.5), 20.0, 38.5, boxstyle="round,pad=0.2,rounding_size=0.5", facecolor="#f0fdf4", edgecolor="#bbf7d0", linewidth=0.8)
    ax.add_patch(c2)
    ax.text(78.6, 38.2, "Section Quality Overview", color="#15803d", fontsize=9.6, fontweight="bold")
    t2 = (
        f"• Section Total Logged: {total_day_mcs} MCs\n"
        f"  - {high_rej_count} Lines > {min_cutoff} pcs (Critical)\n"
        f"  - {total_day_mcs - high_rej_count} Lines <= {min_cutoff} pcs (Controlled)\n"
        f"  - Baseline Registry: {total_plant_mcs} MCs\n\n"
        f"• Last Day Output Lost:\n"
        f"  {total_rej_pcs:,} Pcs / {total_rej_ton:.3f} Ton.\n\n"
        f"• MTD Pace (Day 1–{sel_day_num:02d}):\n"
        f"  {curr_as_of_avg_ton:.2f} T/Day\n"
        f"  (vs {prev_as_of_avg_ton:.2f} T/Day in {prev_abbr})."
    )
    ax.text(78.6, 34.6, t2, color="#166534", fontsize=8.2, linespacing=1.45, va="top")

    plt.subplots_adjust(left=0, right=1, top=1, bottom=0)
    buf = io.BytesIO()
    plt.savefig(buf, format="jpg", facecolor=fig.get_facecolor(), edgecolor="none", dpi=220)
    plt.close(fig)
    buf.seek(0)
    return buf.getvalue()


def m2_generate_cause_pareto_jpg(df_cause_mom, sel_date_obj, sel_day_num, section_label, prev_abbr="Aug", curr_abbr="Sep"):
    fig, ax = plt.subplots(figsize=(19, 11.0), dpi=220)
    fig.patch.set_facecolor('#f1f5f9')
    ax.set_facecolor('#f1f5f9')
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis('off')

    month_name = sel_date_obj.strftime("%B")
    year_num = sel_date_obj.year
    span_text = f"{month_name} 01 – {month_name} {sel_day_num:02d}, {year_num}"

    c_curr_pcs = f"{curr_abbr} Pcs"
    c_prev_pcs = f"{prev_abbr} Pcs"
    c_curr_ton = f"{curr_abbr} Ton"
    c_prev_ton = f"{prev_abbr} Ton"

    tot_curr_pcs = df_cause_mom[c_curr_pcs].sum() if c_curr_pcs in df_cause_mom.columns else 0
    tot_prev_pcs = df_cause_mom[c_prev_pcs].sum() if c_prev_pcs in df_cause_mom.columns else 0
    tot_curr_ton = df_cause_mom[c_curr_ton].sum() if c_curr_ton in df_cause_mom.columns else 0.0
    tot_prev_ton = df_cause_mom[c_prev_ton].sum() if c_prev_ton in df_cause_mom.columns else 0.0
    tot_causes = len(df_cause_mom)

    ax.text(1.5, 98.4, "MTD CAUSE-WISE REJECTION PARETO & MoM REPORT", color='#0f172a', fontsize=16.0, fontweight='bold', va='top')
    ax.text(1.5, 95.8, f"Like-for-Like Root Cause Breakdown ({curr_abbr} 01–{sel_day_num:02d} vs {prev_abbr} 01–{sel_day_num:02d})  |  {section_label}", color='#64748b', fontsize=8.8, va='top')
    
    span_badge = patches.FancyBboxPatch((74.0, 94.8), 24.5, 4.2, boxstyle="round,pad=0.2,rounding_size=0.5", facecolor='#1e293b', edgecolor='none')
    ax.add_patch(span_badge)
    ax.text(86.25, 96.9, f"SPAN: {span_text}", color='#ffffff', fontsize=8.2, fontweight='bold', ha='center', va='center')

    kpis = [
        (f"{curr_abbr.upper()} REJECTION", f"{tot_curr_pcs:,} Pcs", f"{tot_curr_ton:.2f} Tons", "#dc2626"),
        (f"{prev_abbr.upper()} BASELINE", f"{tot_prev_pcs:,} Pcs", f"{tot_prev_ton:.2f} Tons", "#2563eb"),
        ("ACTIVE DEFECT MODES", f"{tot_causes} Modes", "Recorded In Span", "#7c3aed"),
        ("NET MoM VARIANCE", f"{tot_curr_pcs - tot_prev_pcs:+,} Pcs", f"{tot_curr_ton - tot_prev_ton:+.2f} Tons", "#f59e0b"),
    ]
    kpi_w, kpi_gap = 23.0, 1.33
    for i, (title, val, sub, col_bar) in enumerate(kpis):
        x0 = 1.5 + i * (kpi_w + kpi_gap)
        card = patches.FancyBboxPatch((x0, 87.0), kpi_w, 7.2, boxstyle="round,pad=0.15,rounding_size=0.5", facecolor='#ffffff', edgecolor='#cbd5e1', linewidth=0.8)
        ax.add_patch(card)
        top_bar = patches.FancyBboxPatch((x0 + 0.1, 93.75), kpi_w - 0.2, 0.45, boxstyle="round,pad=0.03,rounding_size=0.2", facecolor=col_bar, edgecolor='none')
        ax.add_patch(top_bar)
        ax.text(x0 + kpi_w/2, 92.4, title, color='#64748b', fontsize=7.6, fontweight='bold', ha='center')
        ax.text(x0 + kpi_w/2, 89.6, val, color='#0f172a', fontsize=13.0, fontweight='bold', ha='center')
        ax.text(x0 + kpi_w/2, 87.8, sub, color='#94a3b8', fontsize=6.8, ha='center')

    left_card = patches.FancyBboxPatch((1.5, 1.5), 73.5, 84.0, boxstyle="round,pad=0.25,rounding_size=0.8", facecolor='#ffffff', edgecolor='#cbd5e1', linewidth=1)
    ax.add_patch(left_card)
    ax.text(3.5, 83.5, f"CAUSE-WISE DEFECT VOLUME & TONNAGE MoM COMPARISON — {span_text}", color='#0f172a', fontsize=10.5, fontweight='bold')
    ax.text(73.0, 83.5, f"{tot_causes} Defect Modes Benchmarked", color='#64748b', fontsize=7.8, ha='right')

    right_card = patches.FancyBboxPatch((76.5, 1.5), 22.0, 84.0, boxstyle="round,pad=0.25,rounding_size=0.8", facecolor='#ffffff', edgecolor='#cbd5e1', linewidth=1)
    ax.add_patch(right_card)
    ax.text(78.0, 83.5, "PARETO & MoM SHIFT", color='#0f172a', fontsize=10.5, fontweight='bold')

    left_x = 2.6
    tbl_w = 71.3
    tbl_hdr = patches.Rectangle((left_x, 79.5), tbl_w, 2.6, facecolor='#1e293b', edgecolor='none')
    ax.add_patch(tbl_hdr)
    ax.text(left_x + 1.0, 80.8, "#", color='#ffffff', fontsize=6.8, fontweight='bold', va='center')
    ax.text(left_x + 3.0, 80.8, "DEFECT CAUSE MODE", color='#ffffff', fontsize=6.8, fontweight='bold', va='center')
    ax.text(left_x + 28.0, 80.8, f"{curr_abbr.upper()} PCS", color='#ffffff', fontsize=6.8, fontweight='bold', ha='right', va='center')
    ax.text(left_x + 37.0, 80.8, f"{prev_abbr.upper()} PCS", color='#cbd5e1', fontsize=6.8, fontweight='bold', ha='right', va='center')
    ax.text(left_x + 46.5, 80.8, f"{curr_abbr.upper()} TON", color='#ffffff', fontsize=6.8, fontweight='bold', ha='right', va='center')
    ax.text(left_x + 55.5, 80.8, f"{prev_abbr.upper()} TON", color='#cbd5e1', fontsize=6.8, fontweight='bold', ha='right', va='center')
    ax.text(left_x + 69.5, 80.8, "VARIANCE", color='#ffffff', fontsize=6.8, fontweight='bold', ha='center', va='center')

    row_y = 77.0
    row_step = min(4.2, 73.0 / max(1, min(16, len(df_cause_mom))))
    for r_i, (_, r) in enumerate(df_cause_mom.head(16).iterrows()):
        bg_c = '#f8fafc' if r_i % 2 == 1 else '#ffffff'
        row_bg = patches.Rectangle((left_x, row_y - 1.4), tbl_w, row_step, facecolor=bg_c, edgecolor='none')
        ax.add_patch(row_bg)
        ax.plot([left_x, left_x + tbl_w], [row_y - 1.4, row_y - 1.4], color='#e2e8f0', linewidth=0.45)

        ax.text(left_x + 1.0, row_y + 0.35, f"{r_i + 1}", color='#64748b', fontsize=6.6, va='center')
        ax.text(left_x + 3.0, row_y + 0.35, str(r["Cause"])[:22], color='#0f172a', fontsize=6.8, fontweight='bold', va='center')
        ax.text(left_x + 28.0, row_y + 0.35, f"{int(r[c_curr_pcs]):,}", color='#0f172a', fontsize=6.8, ha='right', va='center')
        ax.text(left_x + 37.0, row_y + 0.35, f"{int(r[c_prev_pcs]):,}", color='#64748b', fontsize=6.6, ha='right', va='center')
        ax.text(left_x + 46.5, row_y + 0.35, f"{r[c_curr_ton]:.3f}", color='#0f172a', fontsize=6.8, ha='right', va='center')
        ax.text(left_x + 55.5, row_y + 0.35, f"{r[c_prev_ton]:.3f}", color='#64748b', fontsize=6.6, ha='right', va='center')

        diff_pcs = r["Variance (Pcs)"] if "Variance (Pcs)" in r else (r[c_curr_pcs] - r[c_prev_pcs])
        badge_x = left_x + 69.5
        if diff_pcs > 0:
            badge_bg = "#fee2e2"
            badge_fg = "#b91c1c"
            badge_txt = f"▲ +{int(diff_pcs):,}"
        else:
            badge_bg = "#dcfce7"
            badge_fg = "#15803d"
            badge_txt = f"▼ {int(diff_pcs):,}"

        ax.add_patch(patches.FancyBboxPatch((badge_x - 4.5, row_y - 0.7), 9.0, 2.1, boxstyle="round,pad=0.08,rounding_size=0.3", facecolor=badge_bg, edgecolor="none"))
        ax.text(badge_x, row_y + 0.35, badge_txt, color=badge_fg, fontsize=6.4, fontweight="bold", ha="center", va="center")
        row_y -= row_step

    c1 = patches.FancyBboxPatch((77.5, 42.5), 20.0, 39.0, boxstyle="round,pad=0.2,rounding_size=0.5", facecolor='#fff7f7', edgecolor='#fecaca', linewidth=0.8)
    ax.add_patch(c1)
    ax.text(78.6, 78.8, "Key MoM Movers", color='#b91c1c', fontsize=9.6, fontweight='bold')

    df_cause_mom["Temp_Diff"] = df_cause_mom[c_curr_pcs] - df_cause_mom[c_prev_pcs]
    top_escalating = df_cause_mom.sort_values("Temp_Diff", ascending=False).iloc[0] if not df_cause_mom.empty else None
    top_declining = df_cause_mom.sort_values("Temp_Diff", ascending=True).iloc[0] if not df_cause_mom.empty else None

    esc_str = f"• Escalating Driver:\n  {top_escalating['Cause']}\n  (+{int(top_escalating['Temp_Diff']):,} pcs vs {prev_abbr})\n\n" if top_escalating is not None else ""
    dec_str = f"• Improved Defect Mode:\n  {top_declining['Cause']}\n  ({int(top_declining['Temp_Diff']):,} pcs vs {prev_abbr})\n\n" if top_declining is not None else ""

    t1 = (
        f"{esc_str}"
        f"{dec_str}"
        f"• Quality Directives:\n"
        f"  - Cushion stability check\n"
        f"  - Mold clamping tonnage\n"
        f"  - Nozzle heat profile."
    )
    ax.text(78.6, 75.2, t1, color='#7f1d1d', fontsize=8.0, linespacing=1.45, va='top')

    c2 = patches.FancyBboxPatch((77.5, 2.5), 20.0, 38.5, boxstyle="round,pad=0.2,rounding_size=0.5", facecolor='#f0fdf4', edgecolor='#bbf7d0', linewidth=0.8)
    ax.add_patch(c2)
    ax.text(78.6, 38.2, "Cumulative Span Summary", color='#15803d', fontsize=9.6, fontweight='bold')
    t2 = (
        f"• Comparison Span:\n"
        f"  Day 01 to Day {sel_day_num:02d}\n"
        f"  ({curr_abbr} vs {prev_abbr}).\n\n"
        f"• Volume Variance:\n"
        f"  - {curr_abbr}: {tot_curr_pcs:,} Pcs ({tot_curr_ton:.2f} T)\n"
        f"  - {prev_abbr}: {tot_prev_pcs:,} Pcs ({tot_prev_ton:.2f} T)\n"
        f"  - Net Shift: {tot_curr_pcs - tot_prev_pcs:+,} Pcs.\n\n"
        f"• Process Control Gate:\n"
        f"  Active line inspection on all\n"
        f"  high-volume running molds."
    )
    ax.text(78.6, 34.6, t2, color='#166534', fontsize=8.0, linespacing=1.45, va='top')

    plt.subplots_adjust(left=0, right=1, top=1, bottom=0)
    buf = io.BytesIO()
    plt.savefig(buf, format="jpg", facecolor=fig.get_facecolor(), edgecolor="none", dpi=220)
    plt.close(fig)
    buf.seek(0)
    return buf.getvalue()


def m2_generate_lineman_report_jpg(df_lineman_mom, sel_date_obj, sel_day_num, section_label, prev_abbr="Aug", curr_abbr="Sep"):
    fig, ax = plt.subplots(figsize=(19, 11.0), dpi=220)
    fig.patch.set_facecolor('#f1f5f9')
    ax.set_facecolor('#f1f5f9')
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis('off')

    month_name = sel_date_obj.strftime("%B")
    year_num = sel_date_obj.year
    span_text = f"{month_name} 01 – {month_name} {sel_day_num:02d}, {year_num}"

    c_curr_pcs = f"{curr_abbr} Pcs"
    c_prev_pcs = f"{prev_abbr} Pcs"
    c_curr_ton = f"{curr_abbr} Ton"
    c_prev_ton = f"{prev_abbr} Ton"

    tot_curr_pcs = df_lineman_mom[c_curr_pcs].sum() if c_curr_pcs in df_lineman_mom.columns else 0
    tot_prev_pcs = df_lineman_mom[c_prev_pcs].sum() if c_prev_pcs in df_lineman_mom.columns else 0
    tot_curr_ton = df_lineman_mom[c_curr_ton].sum() if c_curr_ton in df_lineman_mom.columns else 0.0
    tot_prev_ton = df_lineman_mom[c_prev_ton].sum() if c_prev_ton in df_lineman_mom.columns else 0.0
    tot_linemen = len(df_lineman_mom)

    ax.text(1.5, 98.4, "MTD LINEMAN-WISE REJECTION & MoM AUDIT REPORT", color='#0f172a', fontsize=16.0, fontweight='bold', va='top')
    ax.text(1.5, 95.8, f"Operator Shift Performance & Quality Accountability ({curr_abbr} 01–{sel_day_num:02d} vs {prev_abbr} 01–{sel_day_num:02d})  |  {section_label}", color='#64748b', fontsize=8.8, va='top')
    
    span_badge = patches.FancyBboxPatch((74.0, 94.8), 24.5, 4.2, boxstyle="round,pad=0.2,rounding_size=0.5", facecolor='#1e293b', edgecolor='none')
    ax.add_patch(span_badge)
    ax.text(86.25, 96.9, f"SPAN: {span_text}", color='#ffffff', fontsize=8.2, fontweight='bold', ha='center', va='center')

    kpis = [
        (f"{curr_abbr.upper()} OPERATOR SCRAP", f"{tot_curr_pcs:,} Pcs", f"{tot_curr_ton:.2f} Metric Tons", "#dc2626"),
        (f"{prev_abbr.upper()} BENCHMARK", f"{tot_prev_pcs:,} Pcs", f"{tot_prev_ton:.2f} Metric Tons", "#2563eb"),
        ("BENCHMARKED OPERATORS", f"{tot_linemen} Operators", "Core Roster Audited", "#7c3aed"),
        ("NET SCRAP VARIANCE", f"{tot_curr_pcs - tot_prev_pcs:+,} Pcs", f"{tot_curr_ton - tot_prev_ton:+.2f} Tons", "#f59e0b"),
    ]
    kpi_w, kpi_gap = 23.0, 1.33
    for i, (title, val, sub, col_bar) in enumerate(kpis):
        x0 = 1.5 + i * (kpi_w + kpi_gap)
        card = patches.FancyBboxPatch((x0, 87.0), kpi_w, 7.2, boxstyle="round,pad=0.15,rounding_size=0.5", facecolor='#ffffff', edgecolor='#cbd5e1', linewidth=0.8)
        ax.add_patch(card)
        top_bar = patches.FancyBboxPatch((x0 + 0.1, 93.75), kpi_w - 0.2, 0.45, boxstyle="round,pad=0.03,rounding_size=0.2", facecolor=col_bar, edgecolor='none')
        ax.add_patch(top_bar)
        ax.text(x0 + kpi_w/2, 92.4, title, color='#64748b', fontsize=7.6, fontweight='bold', ha='center')
        ax.text(x0 + kpi_w/2, 89.6, val, color='#0f172a', fontsize=13.0, fontweight='bold', ha='center')
        ax.text(x0 + kpi_w/2, 87.8, sub, color='#94a3b8', fontsize=6.8, ha='center')

    left_card = patches.FancyBboxPatch((1.5, 1.5), 73.5, 84.0, boxstyle="round,pad=0.25,rounding_size=0.8", facecolor='#ffffff', edgecolor='#cbd5e1', linewidth=1)
    ax.add_patch(left_card)
    ax.text(3.5, 83.5, f"OPERATOR LIKE-FOR-LIKE SCRAP MATRIX — {span_text}", color='#0f172a', fontsize=10.5, fontweight='bold')
    ax.text(73.0, 83.5, f"{tot_linemen} Core Operators Active", color='#64748b', fontsize=7.8, ha='right')

    right_card = patches.FancyBboxPatch((76.5, 1.5), 22.0, 84.0, boxstyle="round,pad=0.25,rounding_size=0.8", facecolor='#ffffff', edgecolor='#cbd5e1', linewidth=1)
    ax.add_patch(right_card)
    ax.text(78.0, 83.5, "OPERATIONAL AUDIT", color='#0f172a', fontsize=10.5, fontweight='bold')

    left_x = 2.6
    tbl_w = 71.3
    tbl_hdr = patches.Rectangle((left_x, 79.5), tbl_w, 2.6, facecolor='#1e293b', edgecolor='none')
    ax.add_patch(tbl_hdr)
    ax.text(left_x + 1.0, 80.8, "#", color='#ffffff', fontsize=6.8, fontweight='bold', va='center')
    ax.text(left_x + 3.0, 80.8, "LINEMAN PERSONNEL (ADDED BY)", color='#ffffff', fontsize=6.8, fontweight='bold', va='center')
    ax.text(left_x + 31.0, 80.8, f"{curr_abbr.upper()} PCS", color='#ffffff', fontsize=6.8, fontweight='bold', ha='right', va='center')
    ax.text(left_x + 40.0, 80.8, f"{prev_abbr.upper()} PCS", color='#cbd5e1', fontsize=6.8, fontweight='bold', ha='right', va='center')
    ax.text(left_x + 49.0, 80.8, f"{curr_abbr.upper()} TON", color='#ffffff', fontsize=6.8, fontweight='bold', ha='right', va='center')
    ax.text(left_x + 57.5, 80.8, f"{prev_abbr.upper()} TON", color='#cbd5e1', fontsize=6.8, fontweight='bold', ha='right', va='center')
    ax.text(left_x + 69.5, 80.8, "VARIANCE", color='#ffffff', fontsize=6.8, fontweight='bold', ha='center', va='center')

    row_y = 77.0
    row_step = min(4.4, 73.0 / max(1, tot_linemen))
    for r_i, (_, r) in enumerate(df_lineman_mom.iterrows()):
        bg_c = '#f8fafc' if r_i % 2 == 1 else '#ffffff'
        row_bg = patches.Rectangle((left_x, row_y - 1.4), tbl_w, row_step, facecolor=bg_c, edgecolor='none')
        ax.add_patch(row_bg)
        ax.plot([left_x, left_x + tbl_w], [row_y - 1.4, row_y - 1.4], color='#e2e8f0', linewidth=0.45)

        ax.text(left_x + 1.0, row_y + 0.35, f"{r_i + 1}", color='#64748b', fontsize=6.6, va='center')
        ax.text(left_x + 3.0, row_y + 0.35, str(r["Lineman (Added By)"]), color='#0f172a', fontsize=6.8, fontweight='bold', va='center')
        ax.text(left_x + 31.0, row_y + 0.35, f"{int(r[c_curr_pcs]):,}", color='#0f172a', fontsize=6.8, ha='right', va='center')
        ax.text(left_x + 40.0, row_y + 0.35, f"{int(r[c_prev_pcs]):,}", color='#64748b', fontsize=6.6, ha='right', va='center')
        ax.text(left_x + 49.0, row_y + 0.35, f"{r[c_curr_ton]:.3f}", color='#0f172a', fontsize=6.6, ha='right', va='center')
        ax.text(left_x + 57.5, row_y + 0.35, f"{r[c_prev_ton]:.3f}", color='#64748b', fontsize=6.6, ha='right', va='center')

        diff_pcs = r["Variance (Pcs)"] if "Variance (Pcs)" in r else (r[c_curr_pcs] - r[c_prev_pcs])
        badge_x = left_x + 69.5
        if diff_pcs > 0:
            badge_bg = "#fee2e2"
            badge_fg = "#b91c1c"
            badge_txt = f"▲ +{int(diff_pcs):,}"
        else:
            badge_bg = "#dcfce7"
            badge_fg = "#15803d"
            badge_txt = f"▼ {int(diff_pcs):,}"

        ax.add_patch(patches.FancyBboxPatch((badge_x - 4.5, row_y - 0.7), 9.0, 2.1, boxstyle="round,pad=0.08,rounding_size=0.3", facecolor=badge_bg, edgecolor="none"))
        ax.text(badge_x, row_y + 0.35, badge_txt, color=badge_fg, fontsize=6.4, fontweight="bold", ha="center", va="center")
        row_y -= row_step

    c1 = patches.FancyBboxPatch((77.5, 42.5), 20.0, 39.0, boxstyle="round,pad=0.2,rounding_size=0.5", facecolor='#f8fafc', edgecolor='#cbd5e1', linewidth=0.8)
    ax.add_patch(c1)
    ax.text(78.6, 78.8, "Shift Accountability Summary", color='#0f172a', fontsize=9.6, fontweight='bold')
    t1 = (
        f"• Benchmarking Scope:\n"
        f"  Comparing active core team\n"
        f"  across equal operational days\n"
        f"  (Day 01–{sel_day_num:02d}).\n\n"
        f"• Shift Directives:\n"
        f"  Audit lines with positive\n"
        f"  scrap escalation (&Delta; Pcs).\n"
        f"  Ensure purge scrap weighing\n"
        f"  during shift handover."
    )
    ax.text(78.6, 75.2, t1, color='#334155', fontsize=8.0, linespacing=1.45, va='top')

    c2 = patches.FancyBboxPatch((77.5, 2.5), 20.0, 38.5, boxstyle="round,pad=0.2,rounding_size=0.5", facecolor='#eff6ff', edgecolor='#bfdbfe', linewidth=0.8)
    ax.add_patch(c2)
    ax.text(78.6, 38.2, "Logging Coverage Audit", color='#1d4ed8', fontsize=9.6, fontweight='bold')
    t2 = (
        f"• Audit Span:\n"
        f"  {span_text}\n"
        f"  ({sel_day_num} Operational Days).\n\n"
        f"• Shift Output:\n"
        f"  - {tot_curr_pcs:,} Pcs ({tot_curr_ton:.2f} T)\n"
        f"  - Prior: {tot_prev_pcs:,} Pcs\n"
        f"  - Variance: {tot_curr_pcs - tot_prev_pcs:+,} Pcs.\n\n"
        f"• Compliance Protocol:\n"
        f"  Random physical bin audits\n"
        f"  against ERP recorded logs."
    )
    ax.text(78.6, 34.6, t2, color='#1e3a8a', fontsize=8.0, linespacing=1.45, va='top')

    plt.subplots_adjust(left=0, right=1, top=1, bottom=0)
    buf = io.BytesIO()
    plt.savefig(buf, format="jpg", facecolor=fig.get_facecolor(), edgecolor="none", dpi=220)
    plt.close(fig)
    buf.seek(0)
    return buf.getvalue()


def render_scrap_module():
    c_back, c_title, c_act = st.columns([1.5, 3.5, 1.5], vertical_alignment="center")
    with c_back:
        if st.button("Back to Operations Hub", use_container_width=True):
            st.session_state["active_view"] = "hub_home"
            st.rerun()
    with c_title:
        st.markdown(
            """
            <div style="text-align:center;">
                <div style="color: #b91c1c; font-size: 0.72rem; font-weight: 800; letter-spacing: 0.1em; text-transform: uppercase;">Quality Control Division</div>
                <h3 style="margin:0; font-weight:800; color:#0f172a; letter-spacing: -0.01em;">DAILY REJECTION & DEFECT ANALYTICS</h3>
            </div>
            """,
            unsafe_allow_html=True,
        )
    with c_act:
        if "m2_file_bytes" in st.session_state:
            if st.button("Change Excel File", use_container_width=True):
                st.session_state.pop("m2_file_bytes", None)
                st.rerun()

    st.markdown("<div style='margin-bottom: 1.25rem;'></div>", unsafe_allow_html=True)

    if "m2_file_bytes" not in st.session_state:
        st.markdown(
            """
            <div style="background: #f8fafc; border: 1px solid #cbd5e1; border-left: 5px solid #b91c1c; border-radius: 8px; padding: 1.2rem 1.5rem; margin-bottom: 1.5rem;">
                <div style="color: #b91c1c; font-size: 0.75rem; font-weight: 800; letter-spacing: 0.08em; text-transform: uppercase; margin-bottom: 0.35rem;">
                    DATA INGESTION REQUIREMENT
                </div>
                <div style="color: #0f172a; font-size: 1.05rem; font-weight: 700; margin-bottom: 0.25rem;">
                    Operational Date Range Specification
                </div>
                <div style="color: #475569; font-size: 0.9rem; line-height: 1.5;">
                    Please upload the relevant file containing data <b>from the start of the last month to the current month's latest operational date</b>.
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        c_up, _ = st.columns([2, 1])
        with c_up:
            st.markdown(
                """
                <div style="background:#ffffff; padding:1.5rem; border-radius:8px; border:1px solid #e2e8f0; border-top:4px solid #b91c1c;">
                    <h4 style="margin-top:0; color:#0f172a;">Ingest Rejection / Scrap Workbook</h4>
                    <p style="color:#64748b !important; font-size:0.85rem;">Select the Excel workbook containing monthly defect records (e.g. rej.xlsx).</p>
                </div>
                """,
                unsafe_allow_html=True,
            )
            st.markdown("<div style='margin-bottom: 1rem;'></div>", unsafe_allow_html=True)

            uploaded_file = st.file_uploader(
                "Select Excel File (.xlsx, .xls)", type=["xlsx", "xls"], key="m2_uploader"
            )
            if uploaded_file is not None:
                if st.button(
                    "Ingest Rejection Data & Launch", type="primary", use_container_width=True
                ):
                    st.session_state["m2_file_bytes"] = uploaded_file.getvalue()
                    st.rerun()
    else:
        df_prev, df_curr, df_full = m2_parse_workbook(st.session_state["m2_file_bytes"])

        detected_sec = df_curr["Detected_Section"].iloc[0] if "Detected_Section" in df_curr.columns else "RIP> DPL> Plastic-3"
        sec_name, total_plant_mcs, daily_avail_hrs, pos_map, size_counts, unique_sizes = resolve_section_config(df_curr, fallback_name=detected_sec)

        all_dates = sorted(df_curr["DateStr"].unique().tolist())
        section_display_name = sec_name.split(">")[-1].strip()

        st.markdown('<div class="control-bar-card">', unsafe_allow_html=True)
        c_date, c_cut, c_snap, c_excel = st.columns([1.4, 1.0, 1.3, 1.5], gap="small")
        with c_date:
            sel_date_str = st.selectbox(
                f"Operational Date ({section_display_name})",
                all_dates,
                index=len(all_dates) - 1
            )
        with c_cut:
            min_cutoff = st.number_input("Min Cutoff (Pcs)", min_value=1, value=50, step=10)

        sel_date_obj = pd.to_datetime(sel_date_str)
        sel_day_num = sel_date_obj.day
        day_formatted = sel_date_obj.strftime("%B %d")

        curr_abbr = df_curr["DateClean"].dt.strftime("%b").iloc[0] if not df_curr.empty else "Sep"
        prev_abbr = df_prev["DateClean"].dt.strftime("%b").iloc[0] if not df_prev.empty else "Aug"

        df_day = df_curr[df_curr["DateStr"] == sel_date_str].copy()
        df_as_of = df_curr[df_curr["DateClean"].dt.day <= sel_day_num].copy()
        df_day_filtered = m2_compute_daily_rejection(df_day, pos_map=pos_map, min_qty=min_cutoff)

        if not df_prev.empty:
            df_prev_as_of = df_prev[df_prev["DateClean"].dt.day <= sel_day_num].copy()
            prev_as_of_total_ton = float(df_prev_as_of["Weight_Ton"].sum())
            prev_as_of_avg_ton = prev_as_of_total_ton / max(1, sel_day_num)
        else:
            df_prev_as_of = pd.DataFrame()
            prev_as_of_total_ton, prev_as_of_avg_ton = 0.0, 0.0

        curr_as_of_total_ton = float(df_as_of["Weight_Ton"].sum()) if not df_as_of.empty else 0.0
        curr_as_of_avg_ton = curr_as_of_total_ton / max(1, sel_day_num)

        diff_ton = curr_as_of_avg_ton - prev_as_of_avg_ton
        pct_diff = (diff_ton / prev_as_of_avg_ton * 100.0) if prev_as_of_avg_ton > 0 else 0.0

        if diff_ton > 0:
            variance_line_plain = f"⚠️ Variance: We are producing +{diff_ton:.2f} Tons/Day (+{pct_diff:.1f}%) more rejection compared to {prev_abbr} 01–{sel_day_num:02d}."
            variance_line_html = f'<p style="margin: 0 0 0.75rem 0; color: #dc2626; font-size: 0.85rem;">⚠️ <b>Variance:</b> We are producing <b>+{diff_ton:.2f} Tons/Day (+{pct_diff:.1f}%)</b> more rejection compared to {prev_abbr} 01–{sel_day_num:02d}.</p>'
        elif diff_ton < 0:
            variance_line_plain = f"✅ Variance: We are producing {abs(diff_ton):.2f} Tons/Day ({abs(pct_diff):.1f}%) less rejection compared to {prev_abbr} 01–{sel_day_num:02d}."
            variance_line_html = f'<p style="margin: 0 0 0.75rem 0; color: #16a34a; font-size: 0.85rem;">✅ <b>Variance:</b> We are producing <b>{abs(diff_ton):.2f} Tons/Day ({abs(pct_diff):.1f}%)</b> less rejection compared to {prev_abbr} 01–{sel_day_num:02d}.</p>'
        else:
            variance_line_plain = f"ℹ️ Variance: Daily rejection rate is on par with {prev_abbr} 01–{sel_day_num:02d} baseline."
            variance_line_html = f'<p style="margin: 0 0 0.75rem 0; color: #64748b; font-size: 0.85rem;">ℹ️ <b>Variance:</b> Daily rejection rate is on par with {prev_abbr} 01–{sel_day_num:02d} baseline.</p>'

        total_rej_pcs = int(round(df_day["Qty_Pcs"].sum())) if not df_day.empty else 0
        total_rej_ton = float(df_day["Weight_Ton"].sum()) if not df_day.empty else 0.0
        high_rej_count = len(df_day_filtered)
        mc_col = get_col(df_day, ["Machine", "MC SL"], "Machine")
        total_day_mcs = df_day[mc_col].nunique() if (mc_col in df_day.columns and not df_day.empty) else high_rej_count

        df_cause_day = m2_compute_cause_breakdown(df_day)
        df_cause_as_of = m2_compute_cause_breakdown(df_as_of)

        top3_summary_list = []
        cause_col = get_col(df_day, ["Cause", "Causes"], "Cause")
        if not df_day.empty and cause_col in df_day.columns:
            cause_grp = df_day.groupby(cause_col)["Qty_Pcs"].sum()
            top_cause = cause_grp.idxmax() if not cause_grp.empty else "General"
            top_cause_pcs = int(round(cause_grp.max())) if not cause_grp.empty else 0
            top_cause_pct = (top_cause_pcs / total_rej_pcs * 100.0) if total_rej_pcs > 0 else 0.0

            top3_sorted = cause_grp.sort_values(ascending=False).head(3)
            top3_pcs = top3_sorted.sum()
            top3_pct = (top3_pcs / total_rej_pcs * 100.0) if total_rej_pcs > 0 else 0.0
            for c_name, c_qty in top3_sorted.items():
                c_clean = str(c_name).replace("*", "").strip()
                top3_summary_list.append(
                    (
                        c_clean,
                        int(round(c_qty)),
                        (c_qty / total_rej_pcs * 100.0) if total_rej_pcs > 0 else 0.0,
                    )
                )
        else:
            top_cause, top_cause_pcs, top_cause_pct, top3_pct = "General", 0, 0.0, 0.0

        if not df_day.empty and mc_col in df_day.columns:
            mc_wt_grp = df_day.groupby(mc_col)["Weight_Ton"].sum() * 1000.0
            top_wt_mc = mc_wt_grp.idxmax() if not mc_wt_grp.empty else "-"
            top_wt_kg = float(mc_wt_grp.max()) if not mc_wt_grp.empty else 0.0
        else:
            top_wt_mc, top_wt_kg = "-", 0.0

        jpg_bytes_daily = m2_generate_scrap_jpg(
            df_day_filtered,
            sel_date_obj,
            total_rej_pcs,
            total_rej_ton,
            prev_as_of_total_ton,
            prev_as_of_avg_ton,
            curr_as_of_total_ton,
            curr_as_of_avg_ton,
            high_rej_count,
            total_day_mcs,
            top_cause,
            top_cause_pcs,
            top_cause_pct,
            top_wt_mc,
            top_wt_kg,
            top3_pct,
            top3_summary_list,
            section_display_name,
            total_plant_mcs,
            prev_abbr=prev_abbr,
            curr_abbr=curr_abbr,
            sel_day_num=sel_day_num,
            min_cutoff=min_cutoff,
        )

        excel_bytes_filtered = m2_export_rejection_excel(df_day_filtered)

        with c_snap:
            st.markdown("<div style='height: 1.7rem;'></div>", unsafe_allow_html=True)
            st.download_button(
                label="Download 1-Page JPG",
                data=jpg_bytes_daily,
                file_name=f"Daily_Rejection_Report_{section_display_name}_{sel_date_str}.jpg",
                mime="image/jpeg",
                use_container_width=True,
            )

        with c_excel:
            st.markdown("<div style='height: 1.7rem;'></div>", unsafe_allow_html=True)
            st.download_button(
                label=f"Download Excel (>{min_cutoff} Pcs)",
                data=excel_bytes_filtered,
                file_name=f"Rejection_Log_Over{min_cutoff}Pcs_{section_display_name}_{sel_date_str}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
        st.markdown("</div>", unsafe_allow_html=True)

        k1, k2, k3, k4, k5, k6 = st.columns(6)
        k1.markdown(
            f'<div class="kpi-card slate"><div class="kpi-title">{prev_abbr.upper()} 01–{sel_day_num:02d} TOTAL</div><div class="kpi-val">{prev_as_of_total_ton:.2f} T</div><div class="kpi-sub">Prior MTD Total</div></div>',
            unsafe_allow_html=True,
        )
        k2.markdown(
            f'<div class="kpi-card teal"><div class="kpi-title">{prev_abbr.upper()} 01–{sel_day_num:02d} AVG</div><div class="kpi-val">{prev_as_of_avg_ton:.2f} T/D</div><div class="kpi-sub">Prior Daily Baseline</div></div>',
            unsafe_allow_html=True,
        )
        k3.markdown(
            f'<div class="kpi-card blue"><div class="kpi-title">{curr_abbr.upper()} 01–{sel_day_num:02d} TOTAL</div><div class="kpi-val">{curr_as_of_total_ton:.2f} T</div><div class="kpi-sub">Current MTD Total</div></div>',
            unsafe_allow_html=True,
        )
        k4.markdown(
            f'<div class="kpi-card purple"><div class="kpi-title">{curr_abbr.upper()} 01–{sel_day_num:02d} AVG</div><div class="kpi-val">{curr_as_of_avg_ton:.2f} T/D</div><div class="kpi-sub">Current MTD Pace</div></div>',
            unsafe_allow_html=True,
        )
        k5.markdown(
            f'<div class="kpi-card red"><div class="kpi-title">LAST DAY REJECTION</div><div class="kpi-val">{total_rej_ton:.3f} T</div><div class="kpi-sub">{total_rej_pcs:,} Pcs Lost</div></div>',
            unsafe_allow_html=True,
        )
        k6.markdown(
            f'<div class="kpi-card amber"><div class="kpi-title">CRITICAL MC (&gt;{min_cutoff})</div><div class="kpi-val">{high_rej_count}</div><div class="kpi-sub">Of {total_plant_mcs} Active MCs</div></div>',
            unsafe_allow_html=True,
        )

        st.markdown("<div style='margin-bottom: 1.5rem;'></div>", unsafe_allow_html=True)

        col_left, col_right = st.columns([1.55, 0.95], gap="large")
        with col_left:
            st.markdown(
                f"""
                <div class="section-panel-header">
                    <h4 class="section-panel-title">{section_display_name.upper()} MACHINE REJECTION LOG (&gt;{min_cutoff} PCS)</h4>
                    <span style="color: #64748b; font-size: 0.8rem; font-weight: 600;">{day_formatted}</span>
                </div>
                """,
                unsafe_allow_html=True,
            )
            if not df_day_filtered.empty:
                st.dataframe(
                    df_day_filtered[["Position", "Machine", "Causes", "Qty", "Weight (kg)", "Mold"]],
                    use_container_width=True,
                    hide_index=True,
                    height=390,
                )
            else:
                st.success("No machines exceeded the rejection cutoff threshold today.")

        with col_right:
            approval_text = f"""📋 *{section_display_name.upper()} DAILY SCRAP & REJECTION BRIEF*
📅 *Date:* {day_formatted}

Dear Sir,

These are the line records from *{section_display_name}* where rejection exceeded *{min_cutoff} pieces*:

🔹 *Prev. Month ({prev_abbr} 01–{sel_day_num:02d}):* {prev_as_of_total_ton:.2f} Tons ({prev_as_of_avg_ton:.2f} T/Day)
🔹 *Present Month ({curr_abbr} 01–{sel_day_num:02d}):* {curr_as_of_total_ton:.2f} Tons ({curr_as_of_avg_ton:.2f} T/Day)
{variance_line_plain}

📌 *Please grant your approval to send these items for rejection clearance.*"""

            st.markdown(
                """
                <div class="section-panel-header">
                    <h4 class="section-panel-title">EXECUTIVE APPROVAL TEXT</h4>
                    <span style="color: #64748b; font-size: 0.8rem; font-weight: 600;">Handover Clearance</span>
                </div>
                """,
                unsafe_allow_html=True,
            )
            st.markdown(
                f"""<div class="narrative-block">
                    <p style="margin: 0 0 0.5rem 0; font-weight: 800; color: #1e293b;">{section_display_name.upper()} DAILY SCRAP & REJECTION BRIEF</p>
                    <p style="margin: 0 0 0.75rem 0; color: #64748b; font-size: 0.82rem;"><b>Date:</b> {day_formatted}</p>
                    <p style="margin: 0 0 0.5rem 0;"><b>Dear Sir,</b></p>
                    <p>These are the line records from <b>{section_display_name}</b> where rejection exceeded <b>{min_cutoff} pieces</b>:</p>
                    <p style="margin: 0.5rem 0 0.2rem 0;">&bull; <b>Prev. Month ({prev_abbr} 01–{sel_day_num:02d}):</b> {prev_as_of_total_ton:.2f} Tons ({prev_as_of_avg_ton:.2f} T/Day)</p>
                    <p style="margin: 0 0 0.2rem 0;">&bull; <b>Present Month ({curr_abbr} 01–{sel_day_num:02d}):</b> {curr_as_of_total_ton:.2f} Tons ({curr_as_of_avg_ton:.2f} T/Day)</p>
                    {variance_line_html}
                    <p style="margin: 0.75rem 0 0 0; color: #dc2626; font-weight: 700;">Please grant your approval to send these items for rejection clearance.</p>
                </div>""",
                unsafe_allow_html=True,
            )

            with st.expander("Copy Plain Text for Approval / WhatsApp"):
                st.text_area(
                    "Approval Text", value=approval_text, height=180, label_visibility="collapsed"
                )

        st.divider()

        # Section 3: Cause-Wise Defect Analysis
        df_cause_mom_full = m2_compute_cause_mom_comparison(
            df_prev_as_of, df_as_of, prev_abbr, curr_abbr, include_variances=True
        )
        jpg_bytes_cause_pareto = m2_generate_cause_pareto_jpg(
            df_cause_mom_full,
            sel_date_obj,
            sel_day_num,
            section_display_name,
            prev_abbr=prev_abbr,
            curr_abbr=curr_abbr,
        )

        c_cause_hdr, c_cause_btn = st.columns([3, 1.4], vertical_alignment="center")
        with c_cause_hdr:
            st.markdown("#### CAUSE-WISE REJECTION DEFECT ANALYSIS")
        with c_cause_btn:
            st.download_button(
                label="Download MTD Cause Pareto Report (JPG)",
                data=jpg_bytes_cause_pareto,
                file_name=f"MTD_Cause_Pareto_Report_{section_display_name}_{sel_date_str}.jpg",
                mime="image/jpeg",
                use_container_width=True,
            )

        c_tgl_c, c_chk_var_c = st.columns([2.2, 2.0], vertical_alignment="center")
        with c_tgl_c:
            show_cause_mom = st.toggle(
                f"Compare with Prior Month ({curr_abbr} 01–{sel_day_num:02d} vs {prev_abbr} 01–{sel_day_num:02d})",
                value=False,
                key="tgl_cause_mom"
            )
        with c_chk_var_c:
            show_cause_variances = False
            if show_cause_mom:
                show_cause_variances = st.checkbox("Show Variance & Trend Details", value=False, key="chk_cause_var")

        tab_cause_day, tab_cause_asof = st.tabs([
            f"Selected Date ({day_formatted})",
            f"As of Month-to-Date Defect Pareto (Day 1 – {sel_day_num})",
        ])

        with tab_cause_day:
            display_cause_cols = ["Cause", "Rej_Pcs", "Rej_Kg", "Rej_Ton", "Entries_Count", "MC_Count", "% Share"]
            st.dataframe(df_cause_day[display_cause_cols], use_container_width=True, hide_index=True)

        with tab_cause_asof:
            if show_cause_mom:
                st.caption(f"Showing Side-by-Side Root Cause Comparison: **{curr_abbr} 01–{sel_day_num:02d}** vs **{prev_abbr} 01–{sel_day_num:02d}**")
                df_cause_mom_disp = m2_compute_cause_mom_comparison(
                    df_prev_as_of, df_as_of, prev_abbr, curr_abbr, include_variances=show_cause_variances
                )
                st.dataframe(df_cause_mom_disp, use_container_width=True, hide_index=True)
            else:
                display_cause_cols = ["Cause", "Rej_Pcs", "Rej_Kg", "Rej_Ton", "Entries_Count", "MC_Count", "% Share"]
                st.dataframe(df_cause_as_of[display_cause_cols], use_container_width=True, hide_index=True)

        st.divider()

        # Section 4: Lineman & Shift Quality Audit
        c_l_title, c_down_line = st.columns([3, 1.4], vertical_alignment="center")
        with c_l_title:
            st.markdown("#### LINEMAN & SHIFT QUALITY AUDIT (ADDED BY)")

        lineman_col = get_col(df_as_of, ["Added By", "AddedBy", "Lineman"], "Added By")
        all_operators = sorted([str(op).strip() for op in df_as_of[lineman_col].dropna().unique() if str(op).strip()])
        op_counts = df_as_of[lineman_col].value_counts().to_dict()

        if "selected_linemen_roster" not in st.session_state:
            default_core = [op for op in all_operators if op_counts.get(op, 0) >= 30]
            st.session_state["selected_linemen_roster"] = default_core if default_core else all_operators

        c_pop_op, c_tgl_op, c_chk_var_op = st.columns([1.5, 1.8, 1.4], vertical_alignment="center")

        with c_pop_op:
            num_sel_ops = len(st.session_state["selected_linemen_roster"])
            with st.popover(f"Select Operators ({num_sel_ops}/{len(all_operators)} Selected)"):
                st.markdown("##### Filter Benchmark Roster")
                st.caption("Uncheck casual/relieving staff to prevent benchmark dilution:")

                updated_ops = []
                for op_item in all_operators:
                    is_op_checked = op_item in st.session_state["selected_linemen_roster"]
                    op_entry_cnt = op_counts.get(op_item, 0)
                    chk = st.checkbox(f"{op_item} ({op_entry_cnt} entries)", value=is_op_checked, key=f"chk_op_{op_item}")
                    if chk:
                        updated_ops.append(op_item)

                col_apply_a, col_apply_b = st.columns(2)
                with col_apply_a:
                    if st.button("Apply Selection", type="primary", use_container_width=True):
                        st.session_state["selected_linemen_roster"] = updated_ops if updated_ops else all_operators
                        st.rerun()
                with col_apply_b:
                    if st.button("Reset to Core", use_container_width=True):
                        st.session_state["selected_linemen_roster"] = [op for op in all_operators if op_counts.get(op, 0) >= 30]
                        st.rerun()

        with c_tgl_op:
            show_lineman_mom = st.toggle(
                f"Compare Operators ({curr_abbr} 01–{sel_day_num:02d} vs {prev_abbr} 01–{sel_day_num:02d})",
                value=False,
                key="tgl_lineman_mom"
            )

        with c_chk_var_op:
            show_line_variances = False
            if show_lineman_mom:
                show_line_variances = st.checkbox("Show Variance Details", value=False, key="chk_line_var")

        active_ops = st.session_state["selected_linemen_roster"]
        df_as_of_filtered_ops = df_as_of[df_as_of[lineman_col].isin(active_ops)].copy()
        df_day_filtered_ops = df_day[df_day[lineman_col].isin(active_ops)].copy()
        df_prev_filtered_ops = df_prev_as_of[df_prev_as_of[lineman_col].isin(active_ops)].copy() if not df_prev_as_of.empty else pd.DataFrame()

        df_lineman_as_of = m2_compute_lineman_breakdown(df_as_of_filtered_ops)
        df_lineman_mom = m2_compute_lineman_mom_comparison(
            df_prev_as_of, df_as_of, prev_abbr, curr_abbr, active_ops, include_variances=True
        )

        jpg_bytes_lineman = m2_generate_lineman_report_jpg(
            df_lineman_mom,
            sel_date_obj,
            sel_day_num,
            section_display_name,
            prev_abbr=prev_abbr,
            curr_abbr=curr_abbr,
        )

        with c_down_line:
            st.download_button(
                label="Download Lineman Audit (JPG)",
                data=jpg_bytes_lineman,
                file_name=f"MTD_Lineman_Report_{section_display_name}_{sel_date_str}.jpg",
                mime="image/jpeg",
                use_container_width=True,
            )

        tab_line_asof, tab_shift, tab_matrix, tab_ranks = st.tabs([
            f"Senior Operator Overview (Day 1 – {sel_day_num})",
            "Shift A vs Shift B Audit",
            "Defect Cause x Operator Matrix",
            "Top 5 Defect Drivers by Operator",
        ])

        with tab_line_asof:
            if show_lineman_mom:
                st.caption(f"Comparing Side-by-Side Operator Performance: **{curr_abbr} 01–{sel_day_num:02d}** vs **{prev_abbr} 01–{sel_day_num:02d}**")
                df_lineman_mom_disp = m2_compute_lineman_mom_comparison(
                    df_prev_as_of, df_as_of, prev_abbr, curr_abbr, active_ops, include_variances=show_line_variances
                )
                if not df_lineman_mom_disp.empty:
                    st.dataframe(df_lineman_mom_disp, use_container_width=True, hide_index=True)
                else:
                    st.info("No comparative records found.")
            else:
                display_lineman_cols = ["Lineman (Added By)", "Rej_Pcs", "Rej_Kg", "Rej_Ton", "Logged_Entries", "Machines_Covered", "% Pcs Share", "% Ton Share"]
                if not df_lineman_as_of.empty:
                    st.dataframe(df_lineman_as_of[display_lineman_cols], use_container_width=True, hide_index=True)
                else:
                    st.info("No operator entries matching current filter.")

        with tab_shift:
            c_sh_hdr, c_sh_chk = st.columns([3, 1.4], vertical_alignment="center")
            with c_sh_hdr:
                st.caption(f"Shift Output Comparison (Day 1–{sel_day_num:02d}): **{curr_abbr}** vs **{prev_abbr}**")
            with c_sh_chk:
                show_shift_variances = st.checkbox("Show Shift Variances", value=False, key="chk_shift_var")

            df_shift_mom = m2_compute_shift_mom_comparison(
                df_prev_as_of, df_as_of, prev_abbr, curr_abbr, include_variances=show_shift_variances
            )
            if not df_shift_mom.empty:
                st.dataframe(df_shift_mom, use_container_width=True, hide_index=True)
            else:
                st.info("No shift records available.")

        with tab_matrix:
            c_mat_span, _ = st.columns([2.5, 1.5], vertical_alignment="center")
            with c_mat_span:
                matrix_span = st.radio(
                    "Matrix Period Scope:",
                    options=[f"Present Month ({curr_abbr} 01–{sel_day_num:02d})", f"Previous Month ({prev_abbr} 01–{sel_day_num:02d})"],
                    horizontal=True,
                    key="rad_matrix_span",
                )
            
            target_matrix_df = df_as_of_filtered_ops if curr_abbr in matrix_span else df_prev_filtered_ops
            df_op_matrix = m2_compute_operator_defect_pivot(target_matrix_df)
            if not df_op_matrix.empty:
                st.dataframe(df_op_matrix, use_container_width=True, hide_index=True)
            else:
                st.info("No defect matrix available for selected period.")

        with tab_ranks:
            c_rnk_span, _ = st.columns([2.5, 1.5], vertical_alignment="center")
            with c_rnk_span:
                ranks_span = st.radio(
                    "Rankings Period Scope:",
                    options=[f"Present Month ({curr_abbr} 01–{sel_day_num:02d})", f"Previous Month ({prev_abbr} 01–{sel_day_num:02d})"],
                    horizontal=True,
                    key="rad_ranks_span",
                )

            target_ranks_df = df_as_of_filtered_ops if curr_abbr in ranks_span else df_prev_filtered_ops
            df_op_ranks = m2_compute_operator_rankings(target_ranks_df)
            if not df_op_ranks.empty:
                st.dataframe(df_op_ranks, use_container_width=True, hide_index=True)
            else:
                st.info("No ranking breakdown available for selected period.")

        st.divider()

        # Section 5: Month-over-Month Daily Trend (With 3-Way Shift Filter)
        c_tr_title, c_tr_chk = st.columns([2.8, 1.6], vertical_alignment="center")
        with c_tr_title:
            st.markdown("#### MONTH-OVER-MONTH DAILY REJECTION TREND (PIECES & TONNAGE)")
        with c_tr_chk:
            show_trend_variances = st.checkbox("Show Daily Variance Details", value=False, key="chk_trend_var")

        # 3-Parameter Shift Selector: All by default, Day Shift, Night Shift
        trend_shift_filter = st.radio(
            "Filter Daily Trend by Operational Shift:",
            options=["All (24 Hrs)", "Day Shift (Shift A)", "Night Shift (Shift B)"],
            index=0,
            horizontal=True,
            key="rad_trend_shift",
        )

        trend_display = m2_compute_daily_trend_comparison(
            df_prev,
            df_curr,
            prev_abbr=prev_abbr,
            curr_abbr=curr_abbr,
            shift_filter=trend_shift_filter,
            include_variances=show_trend_variances,
        )
        st.dataframe(trend_display, use_container_width=True, hide_index=True)
